"""Tests for QRE-REPORT-CN-001: report localization, strategy intro block,
metadata compose helper, and display-layer number formatting."""

from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from quant_research_engine.data.metadata_compose import (
    compose_symbol_metadata,
    merge_symbol_metadata,
)
from quant_research_engine.orchestration.research_run import (
    ResearchRunRequest,
    _merge_metadata_into_frame,
    _run_request_payload,
    _strategy_identity,
)
from quant_research_engine.portfolio.batch import BatchBacktestResult
from quant_research_engine.portfolio.full_market_report import (
    COLUMN_LABELS,
    _build_strategy_intro,
    build_drawdown_curve,
    generate_html_report,
    is_ratio_column,
    write_full_market_report,
)
from quant_research_engine.strategies.spec import StrategySpec


class FakeStore:
    """Minimal duck-type store for compose_symbol_metadata."""

    def __init__(self, snapshot: pd.DataFrame | None = None):
        self._snapshot = snapshot if snapshot is not None else pd.DataFrame()

    def _load_reference_daily_basic_snapshot(self, ts_codes, *, as_of=None):
        return self._snapshot.copy()


def _make_spec(**overrides) -> StrategySpec:
    base = dict(
        name="测试策略",
        signal_expression="close > open",
        exit_signal_expression=None,
        market_fields=["open", "close"],
        factor_fields=[],
        forward_horizons=[1],
        quantiles=None,
        holding_days=5,
        stop_loss_pct=0.1,
        take_profit_pct=None,
        trailing_stop_pct=None,
        exclude_st=True,
        exclude_suspended=True,
        min_listing_days=60,
        min_amount=1e7,
    )
    base.update(overrides)
    return StrategySpec(**base)


class TestComposeSymbolMetadata:
    def test_fills_board_buckets_and_industry(self):
        snapshot = pd.DataFrame({
            "ts_code": ["000001.SZ", "300001.SZ", "688001.SH"],
            "total_mv": [1e6, 1e7, 1e8],  # 万元: small/mid/large
            "turnover_rate_f": [0.5, 2.0, 7.0],  # low/mid/high
        })
        meta = compose_symbol_metadata(FakeStore(snapshot), ["000001.SZ", "300001.SZ", "688001.SH"], as_of=date(2026, 1, 1))
        assert list(meta.columns) == ["ts_code", "board", "industry_group", "market_cap_bucket", "liquidity_bucket"]
        by_code = meta.set_index("ts_code")
        assert by_code.loc["000001.SZ", "board"] == "main_board"
        assert by_code.loc["300001.SZ", "board"] == "chinext"
        assert by_code.loc["688001.SH", "board"] == "star_market"
        assert by_code.loc["000001.SZ", "market_cap_bucket"] == "small"
        assert by_code.loc["300001.SZ", "market_cap_bucket"] == "mid"
        assert by_code.loc["688001.SH", "market_cap_bucket"] == "large"
        assert by_code.loc["300001.SZ", "liquidity_bucket"] == "mid"
        assert (meta["industry_group"] == "unknown").all()

    def test_empty_snapshot_still_returns_board(self):
        meta = compose_symbol_metadata(FakeStore(), ["000001.SZ"])
        assert meta.loc[0, "board"] == "main_board"
        assert meta.loc[0, "industry_group"] == "unknown"
        assert pd.isna(meta.loc[0, "market_cap_bucket"])

    def test_missing_symbols_appended(self):
        snapshot = pd.DataFrame({"ts_code": ["000001.SZ"], "total_mv": [1e6], "turnover_rate_f": [0.5]})
        meta = compose_symbol_metadata(FakeStore(snapshot), ["000001.SZ", "000002.SZ"])
        assert set(meta["ts_code"]) == {"000001.SZ", "000002.SZ"}
        assert meta.set_index("ts_code").loc["000002.SZ", "board"] == "main_board"


class TestMergeSymbolMetadata:
    def test_merge_fills_frame_columns(self):
        frame = pd.DataFrame({
            "ts_code": ["000001.SZ", "000001.SZ", "000002.SZ"],
            "close": [10.0, 10.5, 20.0],
        })
        meta = pd.DataFrame({
            "ts_code": ["000001.SZ", "000002.SZ"],
            "board": ["main_board", "main_board"],
            "industry_group": ["unknown", "unknown"],
            "market_cap_bucket": ["mid", "small"],
            "liquidity_bucket": ["low", "high"],
        })
        merged = merge_symbol_metadata(frame, meta)
        assert len(merged) == 3  # row count unchanged
        assert merged["board"].notna().all()
        assert merged.loc[merged["ts_code"] == "000002.SZ", "liquidity_bucket"].iloc[0] == "high"

    def test_existing_columns_filled_not_duplicated(self):
        frame = pd.DataFrame({
            "ts_code": ["000001.SZ"],
            "board": [None],
        })
        meta = pd.DataFrame({"ts_code": ["000001.SZ"], "board": ["main_board"]})
        merged = merge_symbol_metadata(frame, meta)
        assert "board" in merged.columns
        assert not any(c.endswith("_x") or c.endswith("_y") for c in merged.columns)
        assert merged["board"].iloc[0] == "main_board"

    def test_merge_metadata_into_frame_only_when_requested(self, tmp_path):
        store = FakeStore(pd.DataFrame({
            "ts_code": ["000001.SZ"], "total_mv": [1e6], "turnover_rate_f": [0.5],
        }))
        frame = pd.DataFrame({"ts_code": ["000001.SZ"], "close": [10.0]})
        request = ResearchRunRequest(
            strategy_path=tmp_path / "s.yaml", start=date(2026, 1, 1), end=date(2026, 2, 1),
            output_dir=tmp_path, metadata_columns=["board", "market_cap_bucket"],
        )
        merged = _merge_metadata_into_frame(store, frame, request)
        assert merged["board"].iloc[0] == "main_board"
        assert merged["market_cap_bucket"].iloc[0] == "small"

        request_no_meta = ResearchRunRequest(
            strategy_path=tmp_path / "s.yaml", start=date(2026, 1, 1), end=date(2026, 2, 1),
            output_dir=tmp_path,
        )
        unchanged = _merge_metadata_into_frame(store, frame, request_no_meta)
        assert "board" not in unchanged.columns


class TestRunRequestPayload:
    def test_strategy_identity_from_catalog_path(self):
        strategy_id, family = _strategy_identity(Path("strategies/catalog/momentum/my_strategy.yaml"))
        assert strategy_id == "my_strategy"
        assert family == "momentum"

    def test_strategy_identity_without_catalog(self):
        strategy_id, family = _strategy_identity(Path("strategies/high_level/x.yaml"))
        assert strategy_id == "x"
        assert family is None

    def test_payload_contains_identity_and_dates(self, tmp_path):
        spec = _make_spec()
        request = ResearchRunRequest(
            strategy_path=Path("strategies/catalog/momentum/mom.yaml"),
            start=date(2025, 1, 1), end=date(2025, 12, 31),
            output_dir=tmp_path, initial_cash=500000,
        )
        payload = _run_request_payload(spec, request, command="backtest-batch")
        req = payload["request"]
        assert req["strategy_id"] == "mom"
        assert req["strategy_name"] == "测试策略"
        assert req["strategy_family"] == "momentum"
        assert req["date_start"] == "2025-01-01"
        assert req["date_end"] == "2025-12-31"
        assert req["stop_loss_pct"] == 0.1
        assert req["take_profit_pct"] is None
        assert req["exclude_st"] is True
        assert req["min_listing_days"] == 60

    def test_strategy_intro_block(self, tmp_path):
        spec = _make_spec()
        request = ResearchRunRequest(
            strategy_path=Path("strategies/catalog/momentum/mom.yaml"),
            start=date(2025, 1, 1), end=date(2025, 12, 31),
            output_dir=tmp_path,
        )
        payload = _run_request_payload(spec, request, command="backtest-batch")
        intro = _build_strategy_intro(payload)
        assert intro["strategy_name"] == "测试策略"
        assert intro["strategy_id"] == "mom"
        assert intro["strategy_family"] == "momentum"
        assert intro["entry_condition"] == "close > open"
        assert intro["exit_condition"] == "持仓 5 日"  # no exit expression -> holding N days
        assert intro["holding_days"] == 5
        assert intro["stop_loss_pct"] == 0.1
        assert intro["take_profit_pct"] is None
        assert intro["execution"] == "信号日 t，t+1 开盘成交"


class TestRatioColumnDetection:
    def test_name_tokens(self):
        assert is_ratio_column("win_rate")
        assert is_ratio_column("total_return")
        assert is_ratio_column("drawdown_pct")
        assert is_ratio_column("excess_return_vs_buy_hold")
        assert is_ratio_column("sharpe")

    def test_count_columns_are_not_ratios(self):
        assert not is_ratio_column("win_count")
        assert not is_ratio_column("drawdown_duration_days")
        assert not is_ratio_column("trade_count")

    def test_value_range_fallback(self):
        assert is_ratio_column("mystery", pd.Series([0.1, -0.2, 0.05]))
        assert not is_ratio_column("price", pd.Series([10.0, 11.0]))
        assert not is_ratio_column("qty", pd.Series([0.0, 1.0]))  # integer-like


class _FakeSingleResult:
    def __init__(self, symbol: str):
        dates = pd.date_range("2026-01-05", periods=4, freq="B")
        self.orders = pd.DataFrame({
            "symbol": [symbol] * 2, "trade_date": dates[:2], "side": ["buy", "sell"],
            "quantity": [100, 100], "price": [10.0, 11.0], "notional": [1000.0, 1100.0],
            "cost": [5.0, 5.5],
        })
        self.trades = pd.DataFrame({
            "symbol": [symbol], "entry_date": [dates[0]], "exit_date": [dates[1]],
            "quantity": [100], "entry_price": [10.0], "exit_price": [11.0],
            "gross_pnl": [100.0], "cost": [10.0], "net_pnl": [90.0], "exit_reason": ["signal"],
        })
        self.equity_curve = pd.DataFrame({
            "trade_date": dates, "cash": [9000.0] * 4, "market_value": [1000.0] * 4,
            "total_equity": [10000.0, 10100.0, 10050.0, 10200.0],
        })
        self.failed_orders = pd.DataFrame()
        self.summary = {"total_return": 0.02, "trade_count": 1, "win_rate": 1.0, "total_cost": 10.0}


def _make_batch_result() -> BatchBacktestResult:
    r1 = _FakeSingleResult("000001.SZ")
    return BatchBacktestResult(
        symbols=["000001.SZ"],
        results={"000001.SZ": r1},
        summary=pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.02, "trade_count": 1,
            "win_rate": 1.0, "total_cost": 10.0, "board": None, "industry": None,
        }]),
        summary_metrics={"symbol_count": 1, "traded_symbol_count": 1},
    )


def _make_run_request() -> dict:
    return {
        "command": "backtest-batch",
        "strategy_path": "strategies/catalog/momentum/mom.yaml",
        "request": {
            "strategy_id": "mom",
            "strategy_name": "测试策略",
            "strategy_family": "momentum",
            "initial_cash": 1000000,
            "entry_signal": "close > open",
            "exit_signal": None,
            "holding_days": 5,
            "stop_loss_pct": 0.1,
            "take_profit_pct": None,
            "trailing_stop_pct": None,
            "exclude_st": True,
            "exclude_suspended": False,
            "min_listing_days": 60,
            "min_amount": 1e7,
            "date_start": "2025-01-01",
            "date_end": "2025-12-31",
        },
    }


class TestLocalizedReport:
    def test_overview_key_fields_not_none(self, tmp_path):
        paths = write_full_market_report(
            _make_batch_result(), tmp_path, run_request=_make_run_request(), run_id="r1",
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(paths["backtest_report"]))
        ws = wb["Overview"]
        data = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
        assert data.get("策略名称") == "测试策略"
        assert data.get("策略ID") == "mom"
        assert data.get("策略族") == "momentum"
        assert data.get("开始日期") == "2025-01-01"
        assert data.get("结束日期") == "2025-12-31"
        assert data.get("运行ID") == "r1"

    def test_strategy_intro_block_in_overview_sheet(self, tmp_path):
        paths = write_full_market_report(
            _make_batch_result(), tmp_path, run_request=_make_run_request(), run_id="r1",
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(paths["backtest_report"]))
        ws = wb["Overview"]
        keys = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert "策略说明" in keys
        # intro appears before the regular overview keys
        assert keys.index("策略说明") < keys.index("引擎模式")
        data = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
        assert data.get("入场条件") == "close > open"
        assert data.get("出场条件") == "持仓 5 日"
        assert data.get("持仓天数") == 5
        assert data.get("止损比例") == 0.1
        assert data.get("止盈比例") == "-"  # None -> '-', never 'nan'
        assert data.get("剔除ST") is True
        assert data.get("最小上市天数") == 60
        assert data.get("初始资金") == 1000000
        assert data.get("执行方式") == "信号日 t，t+1 开盘成交"

    def test_xlsx_headers_localized_and_number_formats(self, tmp_path):
        paths = write_full_market_report(
            _make_batch_result(), tmp_path, run_request=_make_run_request(), run_id="r1",
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(paths["backtest_report"]))
        ws = wb["SymbolPerformance"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "代码" in headers
        assert "胜率" in headers
        assert "symbol" not in headers
        win_rate_col = headers.index("胜率") + 1
        trade_count_col = headers.index("交易笔数") + 1
        assert ws.cell(row=2, column=win_rate_col).number_format == "0.00%"
        assert ws.cell(row=2, column=trade_count_col).number_format == "0.00"
        # TradesLedger price columns -> 0.00
        ws_tl = wb["TradesLedger"]
        tl_headers = [ws_tl.cell(row=1, column=c).value for c in range(1, ws_tl.max_column + 1)]
        assert "成交价格" in tl_headers
        price_col = tl_headers.index("成交价格") + 1
        assert ws_tl.cell(row=2, column=price_col).number_format == "0.00"

    def test_csv_headers_stay_english(self, tmp_path):
        paths = write_full_market_report(
            _make_batch_result(), tmp_path, run_request=_make_run_request(), run_id="r1",
        )
        for key, first_cols in {
            "trades_ledger": ["run_id", "strategy_id", "symbol", "trade_date"],
            "round_trips": ["run_id", "strategy_id", "symbol", "entry_date"],
            "symbol_performance": ["symbol", "board", "industry", "trade_count"],
        }.items():
            header = paths[key].read_text(encoding="utf-8").splitlines()[0].split(",")
            assert header[: len(first_cols)] == first_cols

    def test_html_contains_chinese_labels_and_intro(self, tmp_path):
        paths = write_full_market_report(
            _make_batch_result(), tmp_path, run_request=_make_run_request(), run_id="r1",
        )
        html = paths["html_report"].read_text(encoding="utf-8")
        assert "回测报告" in html
        assert "策略说明" in html
        assert "入场条件" in html
        assert "出场条件" in html
        assert "close &gt; open" in html
        assert "绩效汇总" in html
        assert "代码" in html  # df table header localized
        assert "0.02" not in html or "2.00%" in html  # ratio shown as percent
        import re
        visible = re.sub(r"data:image/png;base64,[A-Za-z0-9+/=]+", "", html)
        assert "nan" not in visible.lower()

    def test_column_labels_cover_report_columns(self):
        from quant_research_engine.portfolio.full_market_report import (
            _TRADES_LEDGER_COLUMNS,
            _ROUND_TRIPS_COLUMNS,
            _SYMBOL_PERFORMANCE_COLUMNS,
            _EQUITY_CURVE_COLUMNS,
            _DRAWDOWN_CURVE_COLUMNS,
        )
        for col in (
            _TRADES_LEDGER_COLUMNS + _ROUND_TRIPS_COLUMNS + _SYMBOL_PERFORMANCE_COLUMNS
            + _EQUITY_CURVE_COLUMNS + _DRAWDOWN_CURVE_COLUMNS
        ):
            assert col in COLUMN_LABELS, f"missing Chinese label for {col}"
