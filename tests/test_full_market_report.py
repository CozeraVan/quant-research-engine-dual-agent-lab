import json

import numpy as np
import pandas as pd
import pytest

from quant_research_engine.portfolio.batch import BatchBacktestResult
from quant_research_engine.portfolio.full_market_report import (
    build_trades_ledger,
    build_round_trips,
    build_symbol_performance,
    build_strategy_performance,
    build_full_market_equity_curve,
    build_drawdown_curve,
    build_trade_analysis,
    build_capital_usage,
    generate_equity_curve_chart,
    generate_drawdown_chart,
    generate_html_report,
    write_full_market_excel,
    write_full_market_report,
)
from quant_research_engine.portfolio.shared import SharedCashBacktestResult
from quant_research_engine.portfolio.single_symbol import SingleSymbolBacktestResult


def _make_single_result(symbol: str, trade_count: int = 2) -> SingleSymbolBacktestResult:
    dates = pd.date_range("2026-01-05", periods=trade_count * 3, freq="B")
    return SingleSymbolBacktestResult(
        orders=pd.DataFrame({
            "symbol": [symbol] * trade_count,
            "trade_date": dates[:trade_count],
            "side": ["buy"] * trade_count,
            "quantity": [100] * trade_count,
            "price": [10.0] * trade_count,
            "notional": [1000.0] * trade_count,
            "cost": [5.0] * trade_count,
        }),
        trades=pd.DataFrame({
            "symbol": [symbol] * trade_count,
            "entry_date": dates[:trade_count],
            "exit_date": dates[trade_count : trade_count * 2],
            "quantity": [100] * trade_count,
            "entry_price": [10.0] * trade_count,
            "exit_price": [11.0] * trade_count,
            "gross_pnl": [100.0] * trade_count,
            "cost": [10.0] * trade_count,
            "net_pnl": [90.0] * trade_count,
            "exit_reason": ["signal"] * trade_count,
        }),
        equity_curve=pd.DataFrame({
            "trade_date": pd.date_range("2026-01-05", periods=21, freq="B"),
            "cash": [9000.0] * 21,
            "position_quantity": [100] * 21,
            "market_value": [1000.0] * 21,
            "total_equity": np.linspace(10000.0, 11800.0, 21),
        }),
        failed_orders=pd.DataFrame(),
        summary={"initial_cash": 10000.0, "final_equity": 11800.0, "total_return": 0.18, "trade_count": trade_count, "win_rate": 1.0, "total_cost": 20.0},
    )


class TestBuildTradesLedger:
    def test_empty_results(self):
        result = build_trades_ledger({})
        assert result.empty
        assert "symbol" in result.columns
        assert "lot_count" in result.columns
        assert "commission" in result.columns

    def test_combines_symbol_orders(self):
        r1 = _make_single_result("000001.SZ")
        r2 = _make_single_result("000002.SZ")
        df = build_trades_ledger({"000001.SZ": r1, "000002.SZ": r2})
        assert len(df) == 4
        assert set(df["symbol"].unique()) == {"000001.SZ", "000002.SZ"}
        assert (df["side"] == "buy").all()
        assert (df["price_raw"] == 10.0).all()
        assert (df["price_executed"] == 10.0).all()
        assert (df["total_cost"] == 5.0).all()
        assert (df["status"] == "filled").all()
        assert df["commission"].isna().all()
        assert df["lot_count"].isna().all()
        assert df["realized_pnl"].isna().all()  # buy rows → null pnl

    def test_sell_rows_have_null_realized_pnl(self):
        r = _make_single_result("000001.SZ")
        r = SingleSymbolBacktestResult(
            orders=pd.DataFrame({
                "symbol": ["000001.SZ", "000001.SZ"],
                "trade_date": pd.date_range("2026-01-05", periods=2, freq="B"),
                "side": ["buy", "sell"],
                "quantity": [100, 100],
                "price": [10.0, 11.0],
                "notional": [1000.0, 1100.0],
                "cost": [5.0, 5.5],
            }),
            trades=r.trades, equity_curve=r.equity_curve,
            failed_orders=r.failed_orders, summary=r.summary,
        )
        df = build_trades_ledger({"000001.SZ": r})
        assert len(df) == 2
        buy_row = df[df["side"] == "buy"].iloc[0]
        sell_row = df[df["side"] == "sell"].iloc[0]
        assert pd.isna(buy_row["realized_pnl"])
        assert pd.isna(sell_row["realized_pnl"])
        assert sell_row["realized_pnl"] != sell_row["total_cost"]


class TestBuildRoundTrips:
    def test_empty_results(self):
        result = build_round_trips({})
        assert result.empty
        assert "entry_notional" in result.columns
        assert "max_favorable_return" in result.columns

    def test_combines_symbol_trades(self):
        r1 = _make_single_result("000001.SZ")
        r2 = _make_single_result("000002.SZ")
        df = build_round_trips({"000001.SZ": r1, "000002.SZ": r2})
        assert len(df) == 4
        assert set(df["symbol"].unique()) == {"000001.SZ", "000002.SZ"}
        assert (df["gross_pnl"] == 100.0).all()
        assert (df["net_pnl"] == 90.0).all()
        assert (df["entry_price"] == 10.0).all()
        assert (df["exit_price"] == 11.0).all()
        assert df["holding_days"].notna().all()
        assert df["entry_notional"].notna().all()
        assert df["exit_notional"].notna().all()
        # MFE/MAE are computed from the equity segment (entry 10000 -> exit 10180)
        assert (df["max_favorable_return"] > 0).all()
        assert (df["max_adverse_return"] <= 0).all()
        assert df.loc[0, "max_favorable_return"] == pytest.approx(0.018)
        assert df.loc[0, "max_adverse_return"] == pytest.approx(0.0)

    def test_gross_net_return_computed(self):
        r = _make_single_result("000001.SZ")
        df = build_round_trips({"000001.SZ": r})
        assert len(df) == 2
        assert df["gross_return"].iloc[0] == pytest.approx(0.1)  # 100 / 1000
        assert df["net_return"].iloc[0] == pytest.approx(0.09)  # 90 / 1000


class TestBuildSymbolPerformance:
    def test_empty_summary(self):
        df = build_symbol_performance(pd.DataFrame(), {})
        assert df.empty
        assert "symbol" in df.columns

    def test_derives_sharpe_and_drawdown(self):
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0, "win_count": 2,
            "loss_count": 0, "board": "main_board", "industry": "finance",
        }])
        df = build_symbol_performance(summary, {"000001.SZ": r})
        assert len(df) == 1
        row = df.iloc[0]
        assert row["symbol"] == "000001.SZ"
        assert row["board"] == "main_board"
        assert row["strategy_total_return"] == 0.18
        assert row["sharpe"] is not None
        assert row["max_drawdown"] is not None


class TestBuildFullMarketEquityCurve:
    def test_empty_results(self):
        df = build_full_market_equity_curve({})
        assert df.empty

    def test_aggregates_across_symbols(self):
        r1 = _make_single_result("000001.SZ")
        r2 = _make_single_result("000002.SZ")
        df = build_full_market_equity_curve({"000001.SZ": r1, "000002.SZ": r2})
        assert not df.empty
        assert "trade_date" in df.columns
        assert "total_equity" in df.columns
        assert "daily_return" in df.columns
        assert "cumulative_return" in df.columns
        assert "position_count" in df.columns
        assert df["total_equity"].iloc[0] == pytest.approx(20000.0)  # 2 × 10000
        assert df["position_count"].iloc[0] == 2


class TestBuildDrawdownCurve:
    def test_empty_equity(self):
        df = build_drawdown_curve(pd.DataFrame())
        assert df.empty

    def test_computes_running_peak_and_drawdown(self):
        eq = pd.DataFrame({
            "trade_date": pd.date_range("2026-01-05", periods=5, freq="B"),
            "total_equity": [100.0, 110.0, 95.0, 105.0, 120.0],
        })
        df = build_drawdown_curve(eq)
        assert len(df) == 5
        assert df["running_peak"].iloc[2] == 110.0
        assert df["drawdown"].iloc[2] == pytest.approx(-15.0)
        assert df["drawdown_pct"].iloc[2] == pytest.approx(-15.0 / 110.0)
        assert df["drawdown_duration_days"].iloc[2] == 1


class TestBuildStrategyPerformance:
    def test_aggregates_from_round_trips_and_equity(self):
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0,
        }])
        rt = build_round_trips({"000001.SZ": r})
        eq = build_full_market_equity_curve({"000001.SZ": r})
        perf = build_strategy_performance(
            {"symbol_count": 1, "traded_symbol_count": 1},
            summary, rt, eq,
            run_id="test-run", strategy_id="strat-1",
        )
        assert perf["run_id"] == "test-run"
        assert perf["total_trade_count"] == 2
        assert perf["winning_trade_count"] == 2
        assert perf["win_rate"] == 1.0
        assert perf["net_profit"] == pytest.approx(180.0)
        assert perf["sharpe"] is not None
        assert perf["largest_winning_trade"] == pytest.approx(90.0)
        assert "best_symbols" in perf
        assert "worst_symbols" in perf

    def test_benchmark_fields_passthrough(self):
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0,
        }])
        rt = build_round_trips({"000001.SZ": r})
        eq = build_full_market_equity_curve({"000001.SZ": r})
        perf = build_strategy_performance(
            {
                "symbol_count": 1,
                "traded_symbol_count": 1,
                "benchmark_code": "000300.SH",
                "benchmark_aligned_period_count": 3,
                "benchmark_portfolio_total_return": 0.10,
                "benchmark_benchmark_total_return": 0.08,
                "benchmark_excess_total_return": 0.02,
                "benchmark_information_ratio": 0.5,
            },
            summary, rt, eq,
            run_id="test-run", strategy_id="strat-1",
        )
        assert perf["benchmark_code"] == "000300.SH"
        assert perf["benchmark_aligned_period_count"] == 3
        assert perf["benchmark_portfolio_total_return"] == 0.10
        assert perf["benchmark_benchmark_total_return"] == 0.08
        assert perf["benchmark_excess_total_return"] == 0.02
        assert perf["benchmark_information_ratio"] == 0.5

    def test_no_benchmark_fields_when_not_provided(self):
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0,
        }])
        rt = build_round_trips({"000001.SZ": r})
        eq = build_full_market_equity_curve({"000001.SZ": r})
        perf = build_strategy_performance(
            {"symbol_count": 1, "traded_symbol_count": 1},
            summary, rt, eq,
            run_id="test-run", strategy_id="strat-1",
        )
        benchmark_keys = [k for k in perf if k.startswith("benchmark_")]
        assert len(benchmark_keys) == 0

    def test_date_and_excess_fallback_when_summary_metrics_missing(self):
        """QRE-ENG-046：summary_metrics 无 date_start/date_end/excess 时，
        日期取净值曲线首末交易日，超额 = 组合总收益 − 买入持有。"""
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0,
        }])
        rt = build_round_trips({"000001.SZ": r})
        eq = build_full_market_equity_curve({"000001.SZ": r})
        perf = build_strategy_performance(
            {"symbol_count": 1, "traded_symbol_count": 1, "buy_hold_return": 0.10},
            summary, rt, eq,
            run_id="test-run", strategy_id="strat-1",
        )
        eq_first = str(eq["trade_date"].iloc[0])[:10]
        eq_last = str(eq["trade_date"].iloc[-1])[:10]
        assert perf["date_start"] == eq_first
        assert perf["date_end"] == eq_last
        assert perf["excess_return_vs_buy_hold"] == pytest.approx(
            perf["strategy_total_return"] - 0.10
        )

    def test_summary_metrics_values_take_precedence(self):
        """summary_metrics 显式提供日期/超额时优先使用，不走兜底。"""
        r = _make_single_result("000001.SZ")
        summary = pd.DataFrame([{
            "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
            "win_rate": 1.0, "total_cost": 20.0,
        }])
        rt = build_round_trips({"000001.SZ": r})
        eq = build_full_market_equity_curve({"000001.SZ": r})
        perf = build_strategy_performance(
            {"symbol_count": 1, "traded_symbol_count": 1,
             "date_start": "2020-01-01", "date_end": "2020-12-31",
             "buy_hold_return": 0.10, "excess_return_vs_buy_hold": 0.05},
            summary, rt, eq,
            run_id="test-run", strategy_id="strat-1",
        )
        assert perf["date_start"] == "2020-01-01"
        assert perf["date_end"] == "2020-12-31"
        assert perf["excess_return_vs_buy_hold"] == 0.05


class TestBuildTradeAnalysis:
    def test_empty_round_trips(self):
        ta = build_trade_analysis(pd.DataFrame())
        assert ta["win_loss_counts"] == {"win": 0, "loss": 0}
        assert ta["percent_profitable"] is None

    def test_with_trades(self):
        r = _make_single_result("000001.SZ")
        rt = build_round_trips({"000001.SZ": r})
        ta = build_trade_analysis(rt)
        assert ta["win_loss_counts"]["win"] == 2
        assert ta["percent_profitable"] == 1.0
        assert ta["average_pnl"] == pytest.approx(90.0)
        assert ta["largest_winning_trade"] == pytest.approx(90.0)


class TestBuildCapitalUsage:
    def test_computes_cagr(self):
        r = _make_single_result("000001.SZ")
        eq = build_full_market_equity_curve({"000001.SZ": r})
        cu = build_capital_usage(eq, {"net_profit": 180.0, "strategy_total_return": 0.18, "max_drawdown": 0.0})
        assert "cagr" in cu
        assert "average_margin_used" in cu
        assert cu["average_margin_used"] is None
        assert cu["margin_call_count"] == 0


class TestGenerateEquityCurveChart:
    def test_generates_png_for_valid_equity(self, tmp_path):
        equity = pd.DataFrame({
            "trade_date": pd.date_range("2026-01-05", periods=10, freq="B"),
            "total_equity": np.linspace(10000.0, 15000.0, 10),
            "cumulative_return": np.linspace(0.0, 0.5, 10),
        })
        path = generate_equity_curve_chart(equity, tmp_path / "eq.png")
        assert path is not None
        assert path.exists()
        assert path.stat().st_size > 0
        assert path.suffix == ".png"

    def test_returns_none_for_empty_dataframe(self, tmp_path):
        path = generate_equity_curve_chart(pd.DataFrame(), tmp_path / "eq.png")
        assert path is None

    def test_returns_none_for_missing_columns(self, tmp_path):
        equity = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=5)})
        path = generate_equity_curve_chart(equity, tmp_path / "eq.png")
        assert path is None


class TestGenerateDrawdownChart:
    def test_generates_png_for_valid_drawdown(self, tmp_path):
        drawdown = pd.DataFrame({
            "trade_date": pd.date_range("2026-01-05", periods=10, freq="B"),
            "drawdown_pct": [-0.01, -0.02, -0.015, 0.0, -0.005, -0.03, -0.025, -0.01, 0.0, -0.005],
        })
        path = generate_drawdown_chart(drawdown, tmp_path / "dd.png")
        assert path is not None
        assert path.exists()
        assert path.stat().st_size > 0
        assert path.suffix == ".png"

    def test_returns_none_for_empty_dataframe(self, tmp_path):
        path = generate_drawdown_chart(pd.DataFrame(), tmp_path / "dd.png")
        assert path is None

    def test_returns_none_for_missing_columns(self, tmp_path):
        drawdown = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=5)})
        path = generate_drawdown_chart(drawdown, tmp_path / "dd.png")
        assert path is None


class TestGenerateHtmlReport:
    def test_generates_self_contained_html(self, tmp_path):
        tl = pd.DataFrame({"symbol": ["000001.SZ"], "trade_date": ["2026-01-06"], "side": ["buy"], "quantity": [100]})
        rt = pd.DataFrame({"symbol": ["000001.SZ"], "entry_date": ["2026-01-06"], "exit_date": ["2026-01-10"], "side": ["long"], "quantity": [100], "entry_price": [10.0], "exit_price": [11.0], "gross_pnl": [100.0], "net_pnl": [90.0], "return_pct": [10.0], "holding_days": [4]})
        sp = pd.DataFrame({"symbol": ["000001.SZ"], "total_return": [0.10], "trade_count": [1], "win_rate": [1.0], "total_cost": [10.0]})
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100000.0, 100100.0]})
        dd = build_drawdown_curve(eq)
        overview = {"run_id": "test-run", "strategy_name": "test", "engine_mode": "batch", "market": "ashare"}
        strat_perf = {"strategy_total_return": 0.001, "win_rate": 1.0, "sharpe": 0.5, "max_drawdown": 0.0, "total_trade_count": 1, "profit_factor": 1.5}
        ta = {"pnl_distribution_bins": [{"label": ">0%", "count": 1}]}
        cu = {"max_capital_used": 10000.0}

        html = generate_html_report(tmp_path, overview, strat_perf, tl, rt, sp, eq, dd, ta, cu)
        assert "<!DOCTYPE html>" in html
        assert '<html lang="zh-CN">' in html
        assert "</html>" in html
        assert "test-run" in html
        assert "000001.SZ" in html
        # Verify all sections present (Chinese display labels)
        assert "绩效汇总" in html
        assert "图表" in html
        assert "交易分析" in html
        assert "个股表现" in html
        assert "资金使用" in html
        assert "成交流水" in html
        assert "完整回合" in html
        assert "运行概览" in html

    def test_handles_empty_dataframes(self, tmp_path):
        overview = {"run_id": "test-empty"}
        strat_perf = {}
        ta = {}
        cu = {}
        html = generate_html_report(
            tmp_path, overview, strat_perf,
            pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
            pd.DataFrame(), pd.DataFrame(), ta, cu,
        )
        assert "<!DOCTYPE html>" in html
        assert "暂无数据" in html

    def test_embeds_chart_base64_when_provided(self, tmp_path):
        from quant_research_engine.portfolio.full_market_report import generate_equity_curve_chart, generate_drawdown_chart
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=5, freq="B"), "total_equity": [100000.0, 100100.0, 100050.0, 100200.0, 100300.0]})
        dd = build_drawdown_curve(eq)
        ec_path = generate_equity_curve_chart(eq, tmp_path / "ec.png")
        dd_path = generate_drawdown_chart(dd, tmp_path / "dd.png")
        overview = {"run_id": "test-charts"}
        strat_perf = {}
        ta = {}
        cu = {}

        html = generate_html_report(
            tmp_path, overview, strat_perf,
            pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
            eq, dd, ta, cu,
            equity_curve_chart=ec_path,
            drawdown_curve_chart=dd_path,
        )
        assert "data:image/png;base64," in html
        assert '<img src="data:image/png;base64,' in html

    def test_no_chart_tags_when_charts_not_provided(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100000.0, 100100.0]})
        dd = build_drawdown_curve(eq)
        overview = {"run_id": "test-nochart"}
        strat_perf = {}
        ta = {}
        cu = {}

        html = generate_html_report(
            tmp_path, overview, strat_perf,
            pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
            eq, dd, ta, cu,
        )
        assert "暂无净值数据" in html
        assert "暂无回撤数据" in html

    def test_pnl_distribution_renders_bars(self, tmp_path):
        overview = {"run_id": "test-pnl"}
        strat_perf = {}
        ta = {"pnl_distribution_bins": [
            {"label": "<-10%", "count": 5},
            {"label": "-10~0%", "count": 15},
            {"label": "0~10%", "count": 10},
        ]}
        cu = {}
        html = generate_html_report(
            tmp_path, overview, strat_perf,
            pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
            pd.DataFrame(), pd.DataFrame(), ta, cu,
        )
        assert "pnl-bar-row" in html
        assert "pnl-fill" in html


class TestWriteFullMarketExcel:
    def test_creates_workbook_with_all_sheets(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100.0, 110.0]})
        dd = build_drawdown_curve(eq)
        tl = pd.DataFrame({"symbol": ["000001.SZ"], "trade_date": ["2026-01-06"], "side": ["buy"], "quantity": [100]})
        path = write_full_market_excel(
            tmp_path, {}, tl, pd.DataFrame(), pd.DataFrame(),
            eq, dd, {}, {},
        )
        assert path.name == "backtest_report.xlsx"
        assert path.exists()

        from openpyxl import load_workbook
        wb = load_workbook(str(path))
        sheet_names = wb.sheetnames
        assert "Overview" in sheet_names
        assert "TradesLedger" in sheet_names
        assert "RoundTrips" in sheet_names
        assert "SymbolPerformance" in sheet_names
        assert "EquityCurve" in sheet_names
        assert "DrawdownCurve" in sheet_names
        assert "TradeAnalysis" in sheet_names
        assert "CapitalUsage" in sheet_names
        assert "Artifacts" in sheet_names

    def test_embeds_chart_images_when_provided(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=5, freq="B"), "total_equity": np.linspace(100.0, 120.0, 5)})
        dd = build_drawdown_curve(eq)
        chart_path = generate_equity_curve_chart(eq, tmp_path / "eq.png")
        dd_path = generate_drawdown_chart(dd, tmp_path / "dd.png")
        xl_path = write_full_market_excel(
            tmp_path, {}, pd.DataFrame(), pd.DataFrame(),
            pd.DataFrame(), eq, dd, {}, {},
            equity_curve_chart=chart_path, drawdown_curve_chart=dd_path,
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(xl_path))
        assert "EquityCurveChart" in wb.sheetnames
        assert "DrawdownCurveChart" in wb.sheetnames

    def test_skips_chart_sheets_when_not_provided(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100.0, 110.0]})
        dd = build_drawdown_curve(eq)
        xl_path = write_full_market_excel(
            tmp_path, {}, pd.DataFrame(), pd.DataFrame(),
            pd.DataFrame(), eq, dd, {}, {},
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(xl_path))
        assert "EquityCurveChart" not in wb.sheetnames
        assert "DrawdownCurveChart" not in wb.sheetnames

    def test_performance_sheet_populated_from_strategy_perf(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100.0, 110.0]})
        dd = build_drawdown_curve(eq)
        tl = pd.DataFrame({"symbol": ["000001.SZ"], "trade_date": ["2026-01-06"], "side": ["buy"], "quantity": [100]})
        sp = {
            "run_id": "test-run",
            "strategy_id": "strat-1",
            "total_trade_count": 2,
            "sharpe": 1.5,
            "strategy_total_return": 0.15,
            "buy_hold_return": 0.10,
            "excess_return_vs_buy_hold": 0.05,
            "benchmark_code": "000300.SH",
            "benchmark_portfolio_total_return": 0.12,
        }
        path = write_full_market_excel(
            tmp_path, {}, tl, pd.DataFrame(), pd.DataFrame(),
            eq, dd, {}, {},
            strategy_perf=sp,
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(path))
        ws = wb["Performance"]
        perf_data = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}
        assert perf_data.get("运行ID") == "test-run"
        assert perf_data.get("策略ID") == "strat-1"
        assert perf_data.get("交易笔数") == 2
        assert perf_data.get("夏普比率") == 1.5
        assert perf_data.get("策略总收益率") == 0.15
        assert perf_data.get("买入持有收益率") == 0.10
        assert perf_data.get("超额收益率") == 0.05
        assert perf_data.get("基准代码") == "000300.SH"
        assert perf_data.get("benchmark_portfolio_total_return") == 0.12

    def test_performance_sheet_empty_when_none_strategy_perf(self, tmp_path):
        eq = pd.DataFrame({"trade_date": pd.date_range("2026-01-05", periods=2, freq="B"), "total_equity": [100.0, 110.0]})
        dd = build_drawdown_curve(eq)
        tl = pd.DataFrame({"symbol": ["000001.SZ"], "trade_date": ["2026-01-06"], "side": ["buy"], "quantity": [100]})
        path = write_full_market_excel(
            tmp_path, {}, tl, pd.DataFrame(), pd.DataFrame(),
            eq, dd, {}, {},
        )
        from openpyxl import load_workbook
        wb = load_workbook(str(path))
        ws = wb["Performance"]
        assert ws.max_row == 1
        headers = {ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)}
        assert headers == {"项目", "值"}


class TestWriteFullMarketReport:
    def test_writes_all_artifacts(self, tmp_path):
        r1 = _make_single_result("000001.SZ")
        r2 = _make_single_result("000002.SZ")
        result = BatchBacktestResult(
            symbols=["000001.SZ", "000002.SZ"],
            results={"000001.SZ": r1, "000002.SZ": r2},
            summary=pd.DataFrame([
                {"symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2, "win_rate": 1.0, "total_cost": 20.0},
                {"symbol": "000002.SZ", "total_return": 0.12, "trade_count": 2, "win_rate": 0.5, "total_cost": 20.0},
            ]),
            summary_metrics={
                "symbol_count": 2, "traded_symbol_count": 2,
                "positive_return_symbol_rate": 1.0, "mean_total_return": 0.15,
                "median_total_return": 0.15, "total_cost": 40.0,
            },
        )
        paths = write_full_market_report(result, tmp_path, run_id="test-run")

        assert paths["trades_ledger"].name == "trades_ledger.csv"
        assert paths["round_trips"].name == "round_trips.csv"
        assert paths["symbol_performance"].name == "symbol_performance.csv"
        assert paths["strategy_performance"].name == "strategy_performance.json"
        assert paths["equity_curve"].name == "equity_curve.csv"
        assert paths["drawdown_curve"].name == "drawdown_curve.csv"
        assert paths["trade_analysis"].name == "trade_analysis.json"
        assert paths["capital_usage"].name == "capital_usage.json"
        assert paths["backtest_report"].name == "backtest_report.xlsx"
        assert paths["html_report"].name == "backtest_report.html"

        assert pd.read_csv(paths["trades_ledger"]).iloc[0]["symbol"] == "000001.SZ"
        assert pd.read_csv(paths["round_trips"]).iloc[0]["symbol"] == "000001.SZ"
        strat_perf = json.loads(paths["strategy_performance"].read_text(encoding="utf-8"))
        assert strat_perf["run_id"] == "test-run"
        assert strat_perf["total_trade_count"] == 4

    def test_handles_empty_results(self, tmp_path):
        result = BatchBacktestResult(
            symbols=[],
            results={},
            summary=pd.DataFrame(),
            summary_metrics={},
        )
        paths = write_full_market_report(result, tmp_path)
        assert all(p.exists() for p in paths.values())
        assert paths["trades_ledger"].stat().st_size > 0

    def test_generates_chart_artifacts(self, tmp_path):
        r1 = _make_single_result("000001.SZ")
        r2 = _make_single_result("000002.SZ")
        result = BatchBacktestResult(
            symbols=["000001.SZ", "000002.SZ"],
            results={"000001.SZ": r1, "000002.SZ": r2},
            summary=pd.DataFrame([
                {"symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2, "win_rate": 1.0, "total_cost": 20.0},
                {"symbol": "000002.SZ", "total_return": 0.12, "trade_count": 2, "win_rate": 0.5, "total_cost": 20.0},
            ]),
            summary_metrics={
                "symbol_count": 2, "traded_symbol_count": 2,
            },
        )
        paths = write_full_market_report(result, tmp_path)
        assert "equity_curve_chart" in paths
        assert "drawdown_curve_chart" in paths
        assert paths["equity_curve_chart"].exists()
        assert paths["drawdown_curve_chart"].exists()
        assert paths["equity_curve_chart"].suffix == ".png"
        assert paths["drawdown_curve_chart"].suffix == ".png"


class TestBatchResultIntegration:
    def test_write_batch_result_includes_full_market_files(self, tmp_path):
        from quant_research_engine.portfolio.io import write_batch_result

        r = _make_single_result("000001.SZ")
        result = BatchBacktestResult(
            symbols=["000001.SZ"],
            results={"000001.SZ": r},
            summary=pd.DataFrame([{
                "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
                "win_rate": 1.0, "total_cost": 20.0,
            }]),
            summary_metrics={"symbol_count": 1, "traded_symbol_count": 1},
        )
        paths = write_batch_result(result, tmp_path)

        assert "trades_ledger" in paths
        assert "round_trips" in paths
        assert "symbol_performance" in paths
        assert "strategy_performance" in paths
        assert "backtest_report" in paths
        assert paths["backtest_report"].exists()

        manifest = json.loads(paths["report_manifest"].read_text(encoding="utf-8"))
        roles = {entry["role"] for entry in manifest["files"]}
        assert "trades_ledger" in roles
        assert "round_trips" in roles
        assert "symbol_performance" in roles
        assert "strategy_performance" in roles
        assert "full_market_equity_curve" in roles
        assert "drawdown_curve" in roles
        assert "trade_analysis" in roles
        assert "capital_usage" in roles
        assert "full_market_report_workbook" in roles
        assert "html_report" in roles

    def test_existing_roles_still_present(self, tmp_path):
        from quant_research_engine.portfolio.io import write_batch_result

        r = _make_single_result("000001.SZ")
        result = BatchBacktestResult(
            symbols=["000001.SZ"],
            results={"000001.SZ": r},
            summary=pd.DataFrame([{
                "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
                "win_rate": 1.0, "total_cost": 20.0,
            }]),
            summary_metrics={"symbol_count": 1},
        )
        paths = write_batch_result(result, tmp_path)
        manifest = json.loads(paths["report_manifest"].read_text(encoding="utf-8"))
        roles = {entry["role"] for entry in manifest["files"]}
        assert "summary" in roles
        assert "summary_metrics" in roles
        assert "summary_markdown" in roles
        assert "window_summary" in roles
        assert "rolling_summary" in roles
        assert "orders" in roles
        assert "trades" in roles


class TestManifestValidation:
    def test_new_roles_in_required_columns(self):
        from quant_research_engine.portfolio.report_manifest import _REQUIRED_COLUMNS_BY_ROLE
        assert "trades_ledger" in _REQUIRED_COLUMNS_BY_ROLE
        assert "round_trips" in _REQUIRED_COLUMNS_BY_ROLE
        assert "symbol_performance" in _REQUIRED_COLUMNS_BY_ROLE
        assert "full_market_equity_curve" in _REQUIRED_COLUMNS_BY_ROLE
        assert "drawdown_curve" in _REQUIRED_COLUMNS_BY_ROLE
        assert "full_market_report_workbook" in _REQUIRED_COLUMNS_BY_ROLE
        assert "strategy_performance" in _REQUIRED_COLUMNS_BY_ROLE
        assert "trade_analysis" in _REQUIRED_COLUMNS_BY_ROLE
        assert "capital_usage" in _REQUIRED_COLUMNS_BY_ROLE
        assert "html_report" in _REQUIRED_COLUMNS_BY_ROLE

    def test_manifest_validation_passes_for_generated_report(self, tmp_path):
        from quant_research_engine.portfolio.io import write_batch_result
        from quant_research_engine.portfolio.report_manifest import ReportManifest

        r = _make_single_result("000001.SZ")
        result = BatchBacktestResult(
            symbols=["000001.SZ"],
            results={"000001.SZ": r},
            summary=pd.DataFrame([{
                "symbol": "000001.SZ", "total_return": 0.18, "trade_count": 2,
                "win_rate": 1.0, "total_cost": 20.0,
            }]),
            summary_metrics={"symbol_count": 1},
        )
        paths = write_batch_result(result, tmp_path)
        manifest = ReportManifest.from_path(paths["report_manifest"])
        issues = manifest.validate_contents()
        assert issues == []


# ---- shared-cash helpers ----


def _make_shared_result(symbols: list[str] | None = None) -> SharedCashBacktestResult:
    if symbols is None:
        symbols = ["000001.SZ", "000002.SZ"]
    n_orders = len(symbols) * 3
    dates = pd.date_range("2026-01-05", periods=n_orders, freq="B")
    orders = pd.DataFrame({
        "symbol": [s for s in symbols for _ in range(3)],
        "trade_date": dates,
        "side": ["buy"] * n_orders,
        "quantity": [100] * n_orders,
        "price": [10.0] * n_orders,
        "notional": [1000.0] * n_orders,
        "cost": [5.0] * n_orders,
        "reason": ["entry"] * n_orders,
    })
    n_trades = len(symbols) * 2
    trade_dates = pd.date_range("2026-01-05", periods=n_trades * 2, freq="B")
    trades = pd.DataFrame({
        "symbol": [s for s in symbols for _ in range(2)],
        "entry_date": trade_dates[:n_trades],
        "exit_date": trade_dates[n_trades:],
        "quantity": [100] * n_trades,
        "entry_price": [10.0] * n_trades,
        "exit_price": [11.0] * n_trades,
        "gross_pnl": [100.0] * n_trades,
        "cost": [10.0] * n_trades,
        "net_pnl": [90.0] * n_trades,
        "exit_reason": ["signal"] * n_trades,
    })
    eq_dates = pd.date_range("2026-01-05", periods=21, freq="B")
    equity_curve = pd.DataFrame({
        "trade_date": eq_dates,
        "cash": [9000.0] * 21,
        "market_value": [1000.0] * 21,
        "total_equity": np.linspace(10000.0, 11800.0, 21),
        "position_count": [1] * 21,
    })
    return SharedCashBacktestResult(
        orders=orders,
        trades=trades,
        equity_curve=equity_curve,
        failed_orders=pd.DataFrame(),
        summary={
            "initial_cash": 10000.0, "final_equity": 11800.0,
            "total_return": 0.18, "sharpe": 1.2, "max_drawdown": -0.05,
            "max_drawdown_duration_days": 3, "total_cost": 40.0,
            "symbol_count": len(symbols), "traded_symbol_count": len(symbols),
            "trade_count": n_trades, "date_start": "2026-01-05", "date_end": "2026-01-30",
        },
    )


class TestBuildTradesLedgerSharedCash:
    def test_empty_orders(self):
        df = build_trades_ledger(pd.DataFrame(), engine_mode="shared_cash")
        assert df.empty
        assert "symbol" in df.columns

    def test_flat_orders_produce_ledger_rows(self):
        result = _make_shared_result(["000001.SZ"])
        df = build_trades_ledger(result.orders, engine_mode="shared_cash")
        assert len(df) == 3
        assert set(df["symbol"].unique()) == {"000001.SZ"}
        assert (df["side"] == "buy").all()
        assert (df["total_cost"] == 5.0).all()
        assert (df["status"] == "filled").all()

    def test_order_reason_preserved(self):
        result = _make_shared_result(["000001.SZ"])
        df = build_trades_ledger(result.orders, run_id="r1", strategy_id="s1", engine_mode="shared_cash")
        assert df["order_reason"].iloc[0] == "entry"
        assert df["run_id"].iloc[0] == "r1"


class TestBuildRoundTripsSharedCash:
    def test_empty_trades(self):
        df = build_round_trips(pd.DataFrame(), engine_mode="shared_cash")
        assert df.empty
        assert "symbol" in df.columns

    def test_flat_trades_produce_round_trip_rows(self):
        result = _make_shared_result(["000001.SZ", "000002.SZ"])
        df = build_round_trips(result.trades, engine_mode="shared_cash")
        assert len(df) == 4
        assert set(df["symbol"].unique()) == {"000001.SZ", "000002.SZ"}
        assert (df["gross_return"] > 0).all()

    def test_run_id_strategy_id_passthrough(self):
        result = _make_shared_result(["000001.SZ"])
        df = build_round_trips(result.trades, run_id="r99", strategy_id="s99", engine_mode="shared_cash")
        assert df["run_id"].iloc[0] == "r99"
        assert df["strategy_id"].iloc[0] == "s99"


class TestBuildSymbolPerformanceSharedCash:
    def test_empty_round_trips(self):
        df = build_symbol_performance(engine_mode="shared_cash", round_trips=pd.DataFrame(), orders=pd.DataFrame())
        assert df.empty

    def test_per_symbol_aggregation(self):
        result = _make_shared_result(["000001.SZ", "000002.SZ"])
        df = build_symbol_performance(
            engine_mode="shared_cash", round_trips=result.trades, orders=result.orders,
        )
        assert len(df) == 2
        assert set(df["symbol"]) == {"000001.SZ", "000002.SZ"}
        assert (df["trade_count"] == 2).all()
        assert (df["win_count"] == 2).all()
        assert (df["net_profit"] == 180.0).all()
        assert (df["sharpe"].isna()).all()
        assert (df["max_drawdown"].isna()).all()
        assert (df["board"].isna()).all()

    def test_failed_order_count_from_orders(self):
        result = _make_shared_result(["000001.SZ"])
        # Pass empty failed_orders — count should be 0
        df = build_symbol_performance(
            engine_mode="shared_cash", round_trips=result.trades,
            orders=result.orders, failed_orders=pd.DataFrame(),
        )
        assert df["failed_order_count"].iloc[0] == 0

    def test_failed_only_symbol_included(self):
        """零成交但有失败委托的股票也入表（trade_count=0、net_profit=0、
        failed_order_count=N），与 batch 分支“有委托即保留”口径对齐。"""
        result = _make_shared_result(["000001.SZ"])
        failed = pd.DataFrame({
            "symbol": ["000002.SZ", "000002.SZ"],
            "trade_date": pd.date_range("2026-01-05", periods=2, freq="B"),
            "reason": ["insufficient_cash", "insufficient_cash"],
        })
        df = build_symbol_performance(
            engine_mode="shared_cash", round_trips=result.trades,
            orders=result.orders, failed_orders=failed,
        )
        assert set(df["symbol"]) == {"000001.SZ", "000002.SZ"}
        row = df[df["symbol"] == "000002.SZ"].iloc[0]
        assert row["trade_count"] == 0
        assert row["net_profit"] == 0.0
        assert row["failed_order_count"] == 2
        assert row["has_trades"] == False  # noqa: E712
        traded = df[df["symbol"] == "000001.SZ"].iloc[0]
        assert traded["has_trades"] == True  # noqa: E712

    def test_failed_only_symbol_included_when_round_trips_empty(self):
        failed = pd.DataFrame({
            "symbol": ["000003.SZ"],
            "trade_date": ["2026-01-05"],
            "reason": ["suspended"],
        })
        df = build_symbol_performance(
            engine_mode="shared_cash", round_trips=pd.DataFrame(),
            orders=pd.DataFrame(), failed_orders=failed,
        )
        assert len(df) == 1
        row = df.iloc[0]
        assert row["symbol"] == "000003.SZ"
        assert row["trade_count"] == 0
        assert row["failed_order_count"] == 1
        assert row["has_trades"] == False  # noqa: E712

    def test_profit_factor_net_consistent_with_masks(self):
        """回归：毛/净口径混用 bug——一笔 gross_pnl>0 但 net_pnl<0（费用吃掉毛利）
        的交易应被分进亏损桶并以净额计入 gross_loss，避免正值污染 + abs() 失真。"""
        rt = pd.DataFrame([
            # 净盈利 100
            {"symbol": "000001.SZ", "entry_date": "2026-01-05", "exit_date": "2026-01-08",
             "quantity": 100, "entry_price": 10.0, "exit_price": 11.2,
             "gross_pnl": 120.0, "net_pnl": 100.0, "total_cost": 20.0, "exit_reason": "signal"},
            # 毛盈利 10 但费用 30 → 净亏 -20
            {"symbol": "000001.SZ", "entry_date": "2026-02-05", "exit_date": "2026-02-08",
             "quantity": 100, "entry_price": 10.0, "exit_price": 10.1,
             "gross_pnl": 10.0, "net_pnl": -20.0, "total_cost": 30.0, "exit_reason": "signal"},
        ])
        df = build_symbol_performance(engine_mode="shared_cash", round_trips=rt, orders=pd.DataFrame())
        row = df.iloc[0]
        assert row["win_count"] == 1
        assert row["loss_count"] == 1
        # 净口径：盈利桶=100，亏损桶=-20，pf=5.0
        assert row["gross_profit"] == pytest.approx(100.0)
        assert row["gross_loss"] == pytest.approx(-20.0)
        assert row["profit_factor"] == pytest.approx(5.0)
        assert row["net_profit"] == pytest.approx(80.0)


class TestBuildFullMarketEquityCurveSharedCash:
    def test_empty_equity_curve(self):
        df = build_full_market_equity_curve(pd.DataFrame(), engine_mode="shared_cash")
        assert df.empty

    def test_enriches_with_daily_and_cumulative_return(self):
        result = _make_shared_result(["000001.SZ"])
        df = build_full_market_equity_curve(result.equity_curve, engine_mode="shared_cash")
        assert len(df) == 21
        assert "daily_return" in df.columns
        assert "cumulative_return" in df.columns
        assert "turnover" in df.columns
        assert df["daily_return"].notna().sum() > 0
        assert df["cumulative_return"].notna().all()
        assert (df["turnover"].isna()).all()


class TestBuildStrategyPerformanceSharedCash:
    def test_maps_summary_dict_keys(self):
        result = _make_shared_result(["000001.SZ", "000002.SZ"])
        round_trips = build_round_trips(result.trades, engine_mode="shared_cash")
        eq = build_full_market_equity_curve(result.equity_curve, engine_mode="shared_cash")
        perf = build_strategy_performance(
            result.summary, round_trips=round_trips, equity_curve=eq, engine_mode="shared_cash",
        )
        assert perf["sharpe"] == 1.2
        assert perf["max_drawdown"] == -0.05
        assert perf["total_cost"] == 40.0

    def test_null_fields_for_unsupported_batch_metrics(self):
        result = _make_shared_result(["000001.SZ"])
        round_trips = build_round_trips(result.trades, engine_mode="shared_cash")
        eq = build_full_market_equity_curve(result.equity_curve, engine_mode="shared_cash")
        perf = build_strategy_performance(
            result.summary, round_trips=round_trips, equity_curve=eq, engine_mode="shared_cash",
        )
        assert perf["buy_hold_return"] is None
        assert perf["excess_return_vs_buy_hold"] is None
        assert perf["positive_symbol_rate"] is None
        assert perf["mean_symbol_return"] is None
        assert perf["median_symbol_return"] is None

    def test_best_worst_symbols_from_round_trips(self):
        result = _make_shared_result(["000001.SZ", "000002.SZ", "000003.SZ"])
        round_trips = build_round_trips(result.trades, engine_mode="shared_cash")
        eq = build_full_market_equity_curve(result.equity_curve, engine_mode="shared_cash")
        perf = build_strategy_performance(
            result.summary, round_trips=round_trips, equity_curve=eq, engine_mode="shared_cash",
        )
        assert len(perf["best_symbols"]) <= 3
        assert len(perf["worst_symbols"]) <= 3


class TestWriteFullMarketReportSharedCash:
    def test_end_to_end_writes_all_artifacts(self, tmp_path):
        result = _make_shared_result(["000001.SZ", "000002.SZ"])
        paths = write_full_market_report(result, tmp_path, engine_mode="shared_cash")

        assert "trades_ledger" in paths
        assert "round_trips" in paths
        assert "symbol_performance" in paths
        assert "strategy_performance" in paths
        assert "full_market_equity_curve" in paths
        assert "drawdown_curve" in paths
        assert "trade_analysis" in paths
        assert "capital_usage" in paths
        assert "backtest_report" in paths
        assert "html_report" in paths

        assert all(p.exists() for p in paths.values())
        assert pd.read_csv(paths["trades_ledger"]).iloc[0]["symbol"] == "000001.SZ"
        assert pd.read_csv(paths["round_trips"]).iloc[0]["symbol"] == "000001.SZ"
        strat_perf = json.loads(paths["strategy_performance"].read_text(encoding="utf-8"))
        assert strat_perf["total_trade_count"] == 4

    def test_overview_engine_mode_shared_cash(self, tmp_path):
        result = _make_shared_result(["000001.SZ"])
        paths = write_full_market_report(result, tmp_path, engine_mode="shared_cash")
        xlsx_path = paths["backtest_report"]
        from openpyxl import load_workbook
        wb = load_workbook(str(xlsx_path))
        ws = wb["Overview"]
        overview = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}
        assert overview.get("引擎模式") == "shared_cash"
        assert overview.get("初始资金") == 10000.0

    def test_handles_empty_shared_result(self, tmp_path):
        result = SharedCashBacktestResult(
            orders=pd.DataFrame(),
            trades=pd.DataFrame(),
            equity_curve=pd.DataFrame(),
            failed_orders=pd.DataFrame(),
            summary={},
        )
        paths = write_full_market_report(result, tmp_path, engine_mode="shared_cash")
        assert all(p.exists() for p in paths.values())

    def test_generates_chart_artifacts(self, tmp_path):
        result = _make_shared_result(["000001.SZ", "000002.SZ"])
        paths = write_full_market_report(result, tmp_path, engine_mode="shared_cash")
        assert "equity_curve_chart" in paths
        assert "drawdown_curve_chart" in paths
        assert paths["equity_curve_chart"].exists()
        assert paths["drawdown_curve_chart"].exists()
        assert paths["equity_curve_chart"].suffix == ".png"
        assert paths["drawdown_curve_chart"].suffix == ".png"


def test_symbol_performance_filters_zero_trade_symbols():
    from quant_research_engine.portfolio.full_market_report import build_symbol_performance

    summary = pd.DataFrame(
        {
            "symbol": ["A.SZ", "B.SZ", "C.SZ"],
            "trade_count": [2, 0, 0],
            "order_count": [2, 1, 0],
            "total_return": [0.1, 0.0, 0.0],
        }
    )
    result = build_symbol_performance(summary=summary, results={})
    symbols = result["symbol"].tolist() if not result.empty else []
    # A 有成交保留; B 有委托(持仓中)保留; C 无任何交易剔除
    assert symbols == ["A.SZ", "B.SZ"]
    # has_trades 口径：仅 A 有成交；B 仅有委托为 False；summary 缺 failed_order_count 列时兜底 0
    assert result.set_index("symbol")["has_trades"].to_dict() == {"A.SZ": True, "B.SZ": False}
    assert (result["failed_order_count"] == 0).all()


def test_shared_cash_round_trips_cost_breakdown_chain():
    """2026-10-05 费用拆分补齐回归链：shared 引擎交易行拆分列 →
    build_round_trips(shared_cash) 的 commission/stamp_tax/transfer_fee 非空，
    且 三组件和 ≈ total_cost（前端交易磨损卡的直接数据源口径）。"""
    from quant_research_engine.markets.ashare import AShareMarketProfile
    from quant_research_engine.portfolio.shared import (
        SharedCashBacktestEngine, SharedCashBacktestRequest)

    frame = pd.DataFrame({
        "ts_code": ["000001.SZ"] * 5 + ["000002.SZ"] * 5,
        "trade_date": pd.to_datetime(
            ["2026-01-05", "2026-01-06", "2026-01-07", "2026-01-08", "2026-01-09"] * 2),
        "open": [10.0, 10.2, 10.5, 10.8, 11.0, 20.0, 19.8, 19.5, 19.0, 18.8],
        "close": [10.1, 10.4, 10.7, 10.9, 11.1, 19.9, 19.7, 19.3, 18.9, 18.7],
        "rsi_2": [5.0, 20.0, 30.0, 40.0, 50.0, 5.0, 20.0, 30.0, 40.0, 50.0],
        "signal_score": [1.0] * 5 + [3.0] * 5,
    })
    engine = SharedCashBacktestEngine(AShareMarketProfile())
    request = SharedCashBacktestRequest(
        symbols=["000001.SZ", "000002.SZ"],
        initial_cash=100000.0,
        entry_signal="rsi_2 < 10",
        holding_days=2,
        max_positions=2,
    )
    result = engine.run(frame, request)
    trips = build_round_trips(result.trades, run_id="t", strategy_id="s",
                              engine_mode="shared_cash")
    assert not trips.empty
    for col in ("commission", "stamp_tax", "transfer_fee"):
        assert trips[col].notna().all(), col
    for _, r in trips.iterrows():
        assert r["commission"] > 0
        assert r["stamp_tax"] > 0  # 卖出印花税（entry 无印花税，聚合列即 exit 口径）
        assert abs(r["commission"] + r["stamp_tax"] + r["transfer_fee"] - r["total_cost"]) < 1e-6
