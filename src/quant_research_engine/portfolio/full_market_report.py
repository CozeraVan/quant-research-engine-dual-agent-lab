from __future__ import annotations

import base64
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from quant_research_engine.portfolio.risk import summarize_equity_curve

_TRADES_LEDGER_COLUMNS = [
    "run_id", "strategy_id", "symbol", "trade_date", "side", "quantity",
    "lot_count", "price_raw", "price_executed", "notional", "commission",
    "stamp_tax", "transfer_fee", "slippage_amount", "total_cost",
    "realized_pnl", "cash_after", "position_after", "order_reason", "status",
]

_ROUND_TRIPS_COLUMNS = [
    "run_id", "strategy_id", "symbol", "entry_date", "exit_date",
    "holding_days", "entry_price", "exit_price", "quantity",
    "entry_notional", "exit_notional", "gross_pnl", "net_pnl",
    "gross_return", "net_return", "commission", "stamp_tax", "transfer_fee",
    "total_cost", "exit_reason",
    "max_favorable_return", "max_adverse_return",
]

_SYMBOL_PERFORMANCE_COLUMNS = [
    "symbol", "board", "industry", "trade_count", "open_trade_count",
    "win_count", "loss_count", "win_rate", "gross_profit", "gross_loss",
    "net_profit", "profit_factor", "expectancy", "strategy_total_return",
    "buy_hold_return", "excess_return_vs_buy_hold", "sharpe", "sortino",
    "information_ratio_vs_buy_hold", "max_drawdown", "turnover", "total_cost",
    "failed_order_count",
]

_EQUITY_CURVE_COLUMNS = [
    "trade_date", "cash", "market_value", "total_equity",
    "daily_return", "cumulative_return", "position_count", "turnover",
]

_DRAWDOWN_CURVE_COLUMNS = [
    "trade_date", "total_equity", "running_peak", "drawdown",
    "drawdown_pct", "drawdown_duration_days",
]

_EXCEL_SHEET_NAMES = [
    "Overview", "Performance", "TradesLedger", "RoundTrips",
    "SymbolPerformance", "EquityCurve", "DrawdownCurve",
    "TradeAnalysis", "CapitalUsage", "Artifacts",
]

# ---- display-layer localization (xlsx / html only; CSV headers stay English) ----

COLUMN_LABELS: dict[str, str] = {
    # trades ledger
    "run_id": "运行ID",
    "strategy_id": "策略ID",
    "symbol": "代码",
    "trade_date": "交易日期",
    "side": "方向",
    "quantity": "数量",
    "lot_count": "手数",
    "price_raw": "原始价格",
    "price_executed": "成交价格",
    "notional": "成交金额",
    "commission": "佣金",
    "stamp_tax": "印花税",
    "transfer_fee": "过户费",
    "slippage_amount": "滑点金额",
    "total_cost": "总费用",
    "realized_pnl": "已实现盈亏",
    "cash_after": "剩余现金",
    "position_after": "持仓数量",
    "order_reason": "委托原因",
    "status": "状态",
    # round trips
    "entry_date": "入场日期",
    "exit_date": "出场日期",
    "holding_days": "持仓天数",
    "entry_price": "入场价格",
    "exit_price": "出场价格",
    "entry_notional": "入场金额",
    "exit_notional": "出场金额",
    "gross_pnl": "毛盈亏",
    "net_pnl": "净盈亏",
    "gross_return": "毛收益率",
    "net_return": "净收益率",
    "exit_reason": "出场原因",
    "max_favorable_return": "最大浮盈",
    "max_adverse_return": "最大浮亏",
    # symbol performance
    "board": "板块",
    "industry": "行业",
    "trade_count": "交易笔数",
    "open_trade_count": "未平仓笔数",
    "win_count": "盈利笔数",
    "loss_count": "亏损笔数",
    "win_rate": "胜率",
    "gross_profit": "总盈利",
    "gross_loss": "总亏损",
    "net_profit": "净利润",
    "profit_factor": "盈亏比",
    "expectancy": "单笔期望",
    "strategy_total_return": "策略总收益率",
    "buy_hold_return": "买入持有收益率",
    "excess_return_vs_buy_hold": "超额收益率",
    "sharpe": "夏普比率",
    "sortino": "索提诺比率",
    "information_ratio_vs_buy_hold": "信息比率",
    "max_drawdown": "最大回撤",
    "turnover": "换手率",
    "failed_order_count": "失败委托数",
    # equity / drawdown curves
    "cash": "现金",
    "market_value": "持仓市值",
    "total_equity": "总权益",
    "daily_return": "日收益率",
    "cumulative_return": "累计收益率",
    "position_count": "持仓数量",
    "running_peak": "历史峰值",
    "drawdown": "回撤额",
    "drawdown_pct": "回撤幅度",
    "drawdown_duration_days": "回撤持续天数",
}

KEY_LABELS: dict[str, str] = {
    # overview / strategy intro
    "strategy_intro": "策略说明",
    "run_id": "运行ID",
    "holding_days": "持仓天数",
    "strategy_name": "策略名称",
    "strategy_family": "策略族",
    "strategy_description": "策略描述",
    "entry_condition": "入场条件",
    "exit_condition": "出场条件",
    "stop_loss_pct": "止损比例",
    "take_profit_pct": "止盈比例",
    "trailing_stop_pct": "移动止损比例",
    "exclude_st": "剔除ST",
    "exclude_suspended": "剔除停牌",
    "min_listing_days": "最小上市天数",
    "min_amount": "最小成交额",
    "execution": "执行方式",
    "engine_mode": "引擎模式",
    "date_start": "开始日期",
    "date_end": "结束日期",
    "market": "市场",
    "frequency": "频率",
    "universe": "股票池",
    "symbol_count": "股票数量",
    "traded_symbol_count": "实际交易股票数",
    "initial_cash": "初始资金",
    "market_profile": "市场配置",
    "entry_execution": "入场执行",
    "exit_execution": "出场执行",
    "cost_model": "成本模型",
    "slippage_model": "滑点模型",
    "factor_requirements": "因子需求",
    "strategy_source_path": "策略文件路径",
    "generated_at": "生成时间",
    # performance
    "total_trade_count": "交易笔数",
    "winning_trade_count": "盈利笔数",
    "losing_trade_count": "亏损笔数",
    "profit_factor": "盈亏比",
    "expectancy": "单笔期望",
    "recovery_factor": "恢复因子",
    "information_ratio": "信息比率",
    "positive_symbol_rate": "盈利股票占比",
    "mean_symbol_return": "平均个股收益率",
    "median_symbol_return": "个股收益率中位数",
    "max_drawdown_duration_days": "最大回撤持续天数",
    "largest_winning_trade": "最大单笔盈利",
    "largest_losing_trade": "最大单笔亏损",
    "average_winning_trade": "平均盈利",
    "average_losing_trade": "平均亏损",
    "average_win_loss_ratio": "平均盈亏比",
    "max_consecutive_wins": "最大连续盈利次数",
    "max_consecutive_losses": "最大连续亏损次数",
    "average_bars_in_trade": "平均持仓天数",
    "average_bars_in_winning_trades": "盈利交易平均持仓天数",
    "average_bars_in_losing_trades": "亏损交易平均持仓天数",
    "best_symbols": "表现最好股票",
    "worst_symbols": "表现最差股票",
    "benchmark_code": "基准代码",
    # trade analysis
    "pnl_distribution_bins": "盈亏分布",
    "win_loss_counts": "盈亏笔数",
    "percent_profitable": "盈利占比",
    "average_pnl": "平均盈亏",
    "outlier_count": "异常交易数",
    "outlier_pnl": "异常交易盈亏",
    # capital usage
    "cagr": "年化收益率",
    "initial_capital_return": "初始资金收益率",
    "required_account_size": "所需账户规模",
    "account_size_return": "账户规模收益率",
    "net_profit_as_pct_of_max_loss": "净利润/最大亏损",
    "average_margin_used": "平均保证金占用",
    "maximum_margin_used": "最大保证金占用",
    "margin_efficiency": "保证金效率",
    "margin_call_count": "追加保证金次数",
}

_ARTIFACT_DESCRIPTIONS: list[tuple[str, str]] = [
    ("backtest_report.xlsx", "全市场回测报告工作簿"),
    ("trades_ledger.csv", "成交流水（全部股票）"),
    ("round_trips.csv", "完整进出场回合"),
    ("symbol_performance.csv", "分股票绩效汇总"),
    ("strategy_performance.json", "策略整体绩效"),
    ("equity_curve.csv", "全市场净值曲线"),
    ("drawdown_curve.csv", "每日回撤曲线"),
    ("trade_analysis.json", "交易分析（TradingView 风格）"),
    ("capital_usage.json", "资金与保证金占用"),
    ("equity_curve_chart.png", "净值曲线图"),
    ("drawdown_curve_chart.png", "回撤曲线图"),
]

_RATIO_NAME_TOKENS = (
    "return", "rate", "pct", "dd", "win", "excess",
    "drawdown", "sharpe", "volatility",
)
_RATIO_NAME_EXCLUDES = ("count", "day", "quantity", "duration")


def is_ratio_column(name: str, series: pd.Series | None = None) -> bool:
    """Decide whether a column holds a ratio (displayed as 0.00%).

    Name-based: contains a ratio token (return/rate/pct/dd/win/excess/
    drawdown/sharpe/volatility) unless it is clearly a count/duration.
    Fallback: numeric values all within [-1, 1] with at least one
    fractional value.
    """
    lower = str(name).lower()
    if not any(token in lower for token in _RATIO_NAME_EXCLUDES):
        if any(token in lower for token in _RATIO_NAME_TOKENS):
            return True
    if series is not None:
        values = pd.to_numeric(series, errors="coerce").dropna()
        if not values.empty:
            if (values.abs() <= 1.0).all() and (values != values.round(0)).any():
                return True
    return False


def _column_number_format(name: str, series: pd.Series | None = None) -> str:
    return "0.00%" if is_ratio_column(name, series) else "0.00"


def _fmt_display(value: Any, *, ratio: bool = False) -> str:
    """Format one value for the HTML display layer. NaN/None -> '-'."""
    if _is_na(value):
        return "-"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        return f"{float(value):.2%}" if ratio else f"{float(value):.2f}"
    if isinstance(value, dict):
        return json.dumps(_json_ready(value), ensure_ascii=False)
    if isinstance(value, list):
        return json.dumps(_json_ready(value), ensure_ascii=False)[:200]
    return str(value)


def _is_na(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, (float, np.floating)) and pd.isna(value):
        return True
    if value is pd.NA or value is pd.NaT:
        return True
    return False


def _key_label(key: Any) -> str:
    text = str(key)
    return KEY_LABELS.get(text) or COLUMN_LABELS.get(text, text)


def _build_strategy_intro(run_request: dict[str, Any] | None) -> dict[str, Any]:
    """Build the strategy-intro block (策略说明) from the run_request payload."""
    if not run_request:
        return {}
    req = run_request.get("request", {}) or {}
    holding_days = req.get("holding_days")
    exit_condition = req.get("exit_signal") or (f"持仓 {holding_days} 日" if holding_days else None)
    return {
        "strategy_name": req.get("strategy_name"),
        "strategy_id": req.get("strategy_id"),
        "strategy_family": req.get("strategy_family"),
        "entry_condition": req.get("entry_signal"),
        "exit_condition": exit_condition,
        "holding_days": holding_days,
        "stop_loss_pct": req.get("stop_loss_pct"),
        "take_profit_pct": req.get("take_profit_pct"),
        "trailing_stop_pct": req.get("trailing_stop_pct"),
        "exclude_st": req.get("exclude_st"),
        "exclude_suspended": req.get("exclude_suspended"),
        "min_listing_days": req.get("min_listing_days"),
        "min_amount": req.get("min_amount"),
        "initial_cash": req.get("initial_cash"),
        "execution": "信号日 t，t+1 开盘成交",
    }


def build_trades_ledger(
    results: dict[str, Any] | pd.DataFrame,
    run_id: str | None = None,
    strategy_id: str | None = None,
    engine_mode: str = "batch",
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    if engine_mode == "shared_cash":
        orders: pd.DataFrame = results  # type: ignore[assignment]
        if orders.empty:
            return pd.DataFrame(columns=_TRADES_LEDGER_COLUMNS)
        for record in orders.to_dict("records"):
            side = str(record.get("side", ""))
            rows.append({
                "run_id": run_id,
                "strategy_id": strategy_id,
                "symbol": record.get("symbol"),
                "trade_date": _to_date(record.get("trade_date")),
                "side": side,
                "quantity": record.get("quantity"),
                "lot_count": None,
                "price_raw": record.get("price"),
                "price_executed": record.get("price"),
                "notional": record.get("notional"),
                "commission": None,
                "stamp_tax": None,
                "transfer_fee": None,
                "slippage_amount": None,
                "total_cost": record.get("cost"),
                "realized_pnl": None,
                "cash_after": None,
                "position_after": None,
                "order_reason": record.get("reason"),
                "status": "filled",
            })
    else:
        for symbol, single_result in results.items():  # type: ignore[union-attr]
            orders = getattr(single_result, "orders", None)
            if orders is None or (hasattr(orders, "empty") and orders.empty):
                continue
            for record in orders.to_dict("records"):
                side = str(record.get("side", ""))
                price_raw = record.get("price_raw", record.get("price"))
                price_executed = record.get("price")
                quantity = record.get("quantity")
                slippage_amount = (
                    abs(_safe_float(price_executed) - _safe_float(price_raw)) * _safe_float(quantity)
                    if price_raw is not None and price_executed is not None and quantity is not None
                    else None
                )
                rows.append({
                    "run_id": run_id,
                    "strategy_id": strategy_id,
                    "symbol": record.get("symbol", symbol),
                    "trade_date": _to_date(record.get("trade_date")),
                    "side": side,
                    "quantity": quantity,
                    "lot_count": None,
                    "price_raw": price_raw,
                    "price_executed": price_executed,
                    "notional": record.get("notional"),
                    "commission": record.get("commission"),
                    "stamp_tax": record.get("stamp_tax"),
                    "transfer_fee": record.get("transfer_fee"),
                    "slippage_amount": slippage_amount,
                    "total_cost": record.get("cost"),
                    "realized_pnl": None,
                    "cash_after": None,
                    "position_after": None,
                    "order_reason": None,
                    "status": "filled",
                })
    if not rows:
        return pd.DataFrame(columns=_TRADES_LEDGER_COLUMNS)
    return pd.DataFrame(rows, columns=_TRADES_LEDGER_COLUMNS)


def build_round_trips(
    results: dict[str, Any] | pd.DataFrame,
    run_id: str | None = None,
    strategy_id: str | None = None,
    engine_mode: str = "batch",
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    if engine_mode == "shared_cash":
        trades: pd.DataFrame = results  # type: ignore[assignment]
        if trades.empty:
            return pd.DataFrame(columns=_ROUND_TRIPS_COLUMNS)
        for record in trades.to_dict("records"):
            entry_date = _to_date(record.get("entry_date"))
            exit_date = _to_date(record.get("exit_date"))
            holding_days = None
            if entry_date is not None and exit_date is not None:
                try:
                    holding_days = int((pd.Timestamp(exit_date) - pd.Timestamp(entry_date)).days)
                except Exception:
                    holding_days = None
            quantity = record.get("quantity")
            entry_price = record.get("entry_price")
            exit_price = record.get("exit_price")
            entry_notional = _safe_float(quantity) * _safe_float(entry_price) if quantity is not None and entry_price is not None else None
            exit_notional = _safe_float(quantity) * _safe_float(exit_price) if quantity is not None and exit_price is not None else None
            gross_pnl = record.get("gross_pnl")
            net_pnl = record.get("net_pnl")
            gross_return = _safe_float(gross_pnl) / _safe_float(entry_notional) if entry_notional and gross_pnl is not None else None
            net_return = _safe_float(net_pnl) / _safe_float(entry_notional) if entry_notional and net_pnl is not None else None
            rows.append({
                "run_id": run_id,
                "strategy_id": strategy_id,
                "symbol": record.get("symbol"),
                "entry_date": entry_date,
                "exit_date": exit_date,
                "holding_days": holding_days,
                "entry_price": entry_price,
                "exit_price": exit_price,
                "quantity": quantity,
                "entry_notional": entry_notional,
                "exit_notional": exit_notional,
                "gross_pnl": gross_pnl,
                "net_pnl": net_pnl,
                "gross_return": gross_return,
                "net_return": net_return,
                "total_cost": record.get("cost"),
                "exit_reason": record.get("exit_reason"),
                "max_favorable_return": None,
                "max_adverse_return": None,
            })
    else:
        for symbol, single_result in results.items():  # type: ignore[union-attr]
            trades = getattr(single_result, "trades", None)
            if trades is None or (hasattr(trades, "empty") and trades.empty):
                continue
            for record in trades.to_dict("records"):
                entry_date = _to_date(record.get("entry_date"))
                exit_date = _to_date(record.get("exit_date"))
                holding_days = None
                if entry_date is not None and exit_date is not None:
                    try:
                        holding_days = int((pd.Timestamp(exit_date) - pd.Timestamp(entry_date)).days)
                    except Exception:
                        holding_days = None
                quantity = record.get("quantity")
                entry_price = record.get("entry_price")
                exit_price = record.get("exit_price")
                entry_notional = _safe_float(quantity) * _safe_float(entry_price) if quantity is not None and entry_price is not None else None
                exit_notional = _safe_float(quantity) * _safe_float(exit_price) if quantity is not None and exit_price is not None else None
                gross_pnl = record.get("gross_pnl")
                net_pnl = record.get("net_pnl")
                gross_return = _safe_float(gross_pnl) / _safe_float(entry_notional) if entry_notional and gross_pnl is not None else None
                net_return = _safe_float(net_pnl) / _safe_float(entry_notional) if entry_notional and net_pnl is not None else None
                entry_commission = record.get("entry_commission")
                exit_commission = record.get("exit_commission")
                stamp_tax = record.get("exit_stamp_tax")
                transfer_fee = (
                    _safe_float(record.get("entry_transfer_fee") or 0.0)
                    + _safe_float(record.get("exit_transfer_fee") or 0.0)
                    if record.get("entry_transfer_fee") is not None or record.get("exit_transfer_fee") is not None
                    else None
                )
                commission = (
                    _safe_float(entry_commission or 0.0) + _safe_float(exit_commission or 0.0)
                    if entry_commission is not None or exit_commission is not None
                    else None
                )
                rows.append({
                    "run_id": run_id,
                    "strategy_id": strategy_id,
                    "symbol": record.get("symbol", symbol),
                    "entry_date": entry_date,
                    "exit_date": exit_date,
                    "holding_days": holding_days,
                    "entry_price": entry_price,
                    "exit_price": exit_price,
                    "quantity": quantity,
                    "entry_notional": entry_notional,
                    "exit_notional": exit_notional,
                    "gross_pnl": gross_pnl,
                    "net_pnl": net_pnl,
                    "gross_return": gross_return,
                    "net_return": net_return,
                    "commission": commission,
                    "stamp_tax": stamp_tax,
                    "transfer_fee": transfer_fee,
                    "total_cost": record.get("cost"),
                    "exit_reason": record.get("exit_reason"),
                    "max_favorable_return": None,
                    "max_adverse_return": None,
                })
    if not rows:
        return pd.DataFrame(columns=_ROUND_TRIPS_COLUMNS)
    return pd.DataFrame(rows, columns=_ROUND_TRIPS_COLUMNS)


def build_symbol_performance(
    summary: pd.DataFrame | None = None,
    results: dict[str, Any] | None = None,
    engine_mode: str = "batch",
    round_trips: pd.DataFrame | None = None,
    orders: pd.DataFrame | None = None,
    failed_orders: pd.DataFrame | None = None,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    if engine_mode == "shared_cash":
        rt = round_trips if round_trips is not None else pd.DataFrame()
        od = orders if orders is not None else pd.DataFrame()
        fo = failed_orders if failed_orders is not None else pd.DataFrame()
        if rt.empty:
            return pd.DataFrame(columns=_SYMBOL_PERFORMANCE_COLUMNS)
        symbols = sorted(rt["symbol"].dropna().unique())
        failed_by_symbol: dict[str, int] = {}
        if not fo.empty and "symbol" in fo.columns:
            failed_by_symbol = fo.groupby("symbol").size().to_dict()
        for symbol in symbols:
            sym_trades = rt[rt["symbol"] == symbol]
            trade_count = len(sym_trades)
            win_mask = sym_trades["net_pnl"] > 0
            loss_mask = sym_trades["net_pnl"] < 0
            win_count = int(win_mask.sum())
            loss_count = int(loss_mask.sum())
            win_rate = win_count / trade_count if trade_count > 0 else None
            gross_profit = float(sym_trades.loc[win_mask, "gross_pnl"].sum()) if win_count > 0 else 0.0
            gross_loss = float(sym_trades.loc[loss_mask, "gross_pnl"].sum()) if loss_count > 0 else 0.0
            net_profit = float(sym_trades["net_pnl"].sum())
            profit_factor = abs(gross_profit / gross_loss) if gross_loss != 0 else None
            positive_pnl = sym_trades.loc[win_mask, "net_pnl"]
            negative_pnl = sym_trades.loc[loss_mask, "net_pnl"]
            avg_win = float(positive_pnl.mean()) if not positive_pnl.empty else 0.0
            avg_loss = float(negative_pnl.mean()) if not negative_pnl.empty else 0.0
            expectancy = None
            if win_rate is not None and trade_count > 0:
                expectancy = float(win_rate) * avg_win - (1.0 - float(win_rate)) * abs(avg_loss)
            open_trades = od[(od["symbol"] == symbol) & (od.get("reason") == "entry_exit")] if not od.empty else pd.DataFrame()
            open_trade_count = len(open_trades) - (2 * trade_count) if not open_trades.empty else None
            rows.append({
                "symbol": str(symbol),
                "board": None,
                "industry": None,
                "trade_count": trade_count,
                "open_trade_count": open_trade_count,
                "win_count": win_count,
                "loss_count": loss_count,
                "win_rate": win_rate,
                "gross_profit": gross_profit,
                "gross_loss": gross_loss,
                "net_profit": net_profit,
                "profit_factor": profit_factor,
                "expectancy": expectancy,
                "strategy_total_return": None,
                "buy_hold_return": None,
                "excess_return_vs_buy_hold": None,
                "sharpe": None,
                "sortino": None,
                "information_ratio_vs_buy_hold": None,
                "max_drawdown": None,
                "turnover": None,
                "total_cost": float(sym_trades["total_cost"].sum()) if "total_cost" in sym_trades.columns else None,
                "failed_order_count": failed_by_symbol.get(symbol, 0),
            })
    else:
        for record in summary.to_dict("records"):  # type: ignore[union-attr]
            # 只保留存在交易行为(成交或委托)的标的行，无交易标的从个股表现中剔除
            if not (record.get("trade_count") or record.get("order_count")):
                continue
            symbol = str(record.get("symbol", ""))
            single_result = results.get(symbol) if results is not None else None  # type: ignore[union-attr]
            equity_curve = getattr(single_result, "equity_curve", None) if single_result is not None else None
            risk_metrics = {}
            if equity_curve is not None and not equity_curve.empty and "total_equity" in equity_curve.columns:
                risk_metrics = summarize_equity_curve(equity_curve)

            win_count = record.get("win_count")
            loss_count = record.get("loss_count")
            trade_count = record.get("trade_count")
            win_rate = record.get("win_rate")
            net_profit = record.get("net_profit", record.get("total_return"))

            profit_factor = None
            gross_profit = record.get("gross_profit")
            gross_loss = record.get("gross_loss")
            if gross_profit is not None and gross_loss is not None and gross_loss != 0:
                profit_factor = abs(_safe_float(gross_profit) / _safe_float(gross_loss))

            expectancy = None
            if win_rate is not None and trade_count is not None and trade_count > 0:
                avg_win = _safe_float(record.get("avg_win", 0.0))
                avg_loss = _safe_float(record.get("avg_loss", 0.0))
                if avg_win is not None and avg_loss is not None:
                    expectancy = float(win_rate) * avg_win - (1.0 - float(win_rate)) * abs(avg_loss)

            rows.append({
                "symbol": symbol,
                "board": record.get("board"),
                "industry": record.get("industry"),
                "trade_count": trade_count,
                "open_trade_count": record.get("open_trade_count"),
                "win_count": win_count,
                "loss_count": loss_count,
                "win_rate": win_rate,
                "gross_profit": gross_profit,
                "gross_loss": gross_loss,
                "net_profit": net_profit,
                "profit_factor": profit_factor,
                "expectancy": expectancy,
                "strategy_total_return": record.get("total_return"),
                "buy_hold_return": record.get("buy_hold_return"),
                "excess_return_vs_buy_hold": record.get("excess_return_vs_buy_hold"),
                "sharpe": risk_metrics.get("sharpe"),
                "sortino": None,
                "information_ratio_vs_buy_hold": record.get("information_ratio_vs_buy_hold"),
                "max_drawdown": record.get("max_drawdown") or risk_metrics.get("max_drawdown"),
                "turnover": record.get("turnover"),
                "total_cost": record.get("total_cost"),
                "failed_order_count": record.get("failed_order_count", 0),
            })
    if not rows:
        return pd.DataFrame(columns=_SYMBOL_PERFORMANCE_COLUMNS)
    return pd.DataFrame(rows, columns=_SYMBOL_PERFORMANCE_COLUMNS)


def build_full_market_equity_curve(
    results: dict[str, Any] | pd.DataFrame,
    engine_mode: str = "batch",
) -> pd.DataFrame:
    if engine_mode == "shared_cash":
        eq: pd.DataFrame = results.copy()  # type: ignore[assignment, union-attr]
        if eq.empty or "total_equity" not in eq.columns:
            return pd.DataFrame(columns=_EQUITY_CURVE_COLUMNS)
        eq["trade_date"] = pd.to_datetime(eq["trade_date"])
        eq = eq.sort_values("trade_date")
        daily_returns = []
        cumulative = 1.0
        prev_equity = None
        for _, row in eq.iterrows():
            te = _safe_float(row.get("total_equity"))
            dr = None
            if prev_equity is not None and prev_equity > 0:
                dr = (te / prev_equity) - 1.0
            daily_returns.append(dr)
            cumulative *= (1.0 + (dr or 0.0))
            prev_equity = te
        return pd.DataFrame({
            "trade_date": eq["trade_date"],
            "cash": eq.get("cash", 0.0),
            "market_value": eq.get("market_value", 0.0),
            "total_equity": eq["total_equity"],
            "daily_return": daily_returns,
            "cumulative_return": [cumulative - 1.0] * len(eq),
            "position_count": eq.get("position_count", 0),
            "turnover": None,
        }, columns=_EQUITY_CURVE_COLUMNS)

    # batch mode
    all_dates: set[pd.Timestamp] = set()
    symbol_curves: dict[str, pd.DataFrame] = {}
    for symbol, single_result in results.items():  # type: ignore[union-attr]
        eq = getattr(single_result, "equity_curve", None)
        if eq is None or eq.empty or "trade_date" not in eq.columns:
            continue
        eq = eq.copy()
        eq["trade_date"] = pd.to_datetime(eq["trade_date"])
        eq = eq.sort_values("trade_date")
        symbol_curves[symbol] = eq
        all_dates.update(eq["trade_date"].dropna().tolist())

    if not all_dates:
        return pd.DataFrame(columns=_EQUITY_CURVE_COLUMNS)

    dates = sorted(all_dates)
    rows: list[dict[str, Any]] = []
    cumulative_return = 1.0
    prev_total_equity: float | None = None

    for date in dates:
        total_equity = 0.0
        cash = 0.0
        market_value = 0.0
        position_count = 0
        for eq in symbol_curves.values():
            day_data = eq[eq["trade_date"] == date]
            if day_data.empty:
                continue
            total_equity += _safe_float(day_data["total_equity"].iloc[0])
            if "cash" in day_data.columns:
                cash += _safe_float(day_data["cash"].iloc[0])
            if "market_value" in day_data.columns:
                market_value += _safe_float(day_data["market_value"].iloc[0])
            pos_qty = _safe_float(day_data.get("position_quantity", pd.Series([0])).iloc[0])
            if pos_qty is not None and pos_qty > 0:
                position_count += 1
        daily_return = None
        if prev_total_equity is not None and prev_total_equity > 0:
            daily_return = (total_equity / prev_total_equity) - 1.0
        cumulative_return *= (1.0 + (daily_return or 0.0))

        rows.append({
            "trade_date": date,
            "cash": cash,
            "market_value": market_value,
            "total_equity": total_equity,
            "daily_return": daily_return,
            "cumulative_return": cumulative_return - 1.0,
            "position_count": position_count,
            "turnover": None,
        })
        prev_total_equity = total_equity

    return pd.DataFrame(rows, columns=_EQUITY_CURVE_COLUMNS)


def build_drawdown_curve(equity_curve: pd.DataFrame) -> pd.DataFrame:
    if equity_curve.empty or "total_equity" not in equity_curve.columns:
        return pd.DataFrame(columns=_DRAWDOWN_CURVE_COLUMNS)

    eq = equity_curve.copy()
    eq["running_peak"] = eq["total_equity"].cummax()
    eq["drawdown"] = eq["total_equity"] - eq["running_peak"]
    eq["drawdown_pct"] = np.where(
        eq["running_peak"] > 0,
        eq["drawdown"] / eq["running_peak"],
        0.0,
    )

    duration_days: list[int | None] = []
    current_duration = 0
    for _, row in eq.iterrows():
        dd_val = float(row["drawdown"])
        if dd_val < 0:
            current_duration += 1
        elif dd_val == 0:
            current_duration = 0
        else:
            current_duration = 0
        duration_days.append(current_duration if current_duration > 0 else 0)

    return pd.DataFrame({
        "trade_date": eq["trade_date"],
        "total_equity": eq["total_equity"],
        "running_peak": eq["running_peak"],
        "drawdown": eq["drawdown"],
        "drawdown_pct": eq["drawdown_pct"],
        "drawdown_duration_days": duration_days,
    }, columns=_DRAWDOWN_CURVE_COLUMNS)


def build_strategy_performance(
    summary_metrics: dict[str, Any],
    summary: pd.DataFrame | None = None,
    round_trips: pd.DataFrame | None = None,
    equity_curve: pd.DataFrame | None = None,
    run_id: str | None = None,
    strategy_id: str | None = None,
    strategy_name: str | None = None,
    strategy_family: str | None = None,
    engine_mode: str = "batch",
) -> dict[str, Any]:
    rt = round_trips if round_trips is not None else pd.DataFrame()
    eq = equity_curve if equity_curve is not None else pd.DataFrame()

    risk = {}
    if not eq.empty and "total_equity" in eq.columns:
        risk = summarize_equity_curve(eq)

    win_count = int((rt["net_pnl"] > 0).sum()) if not rt.empty else 0
    loss_count = int((rt["net_pnl"] < 0).sum()) if not rt.empty else 0
    total_trades = len(rt)

    positive_pnl = rt[rt["net_pnl"] > 0]["net_pnl"] if not rt.empty else pd.Series(dtype=float)
    negative_pnl = rt[rt["net_pnl"] < 0]["net_pnl"] if not rt.empty else pd.Series(dtype=float)

    sortino = None
    if not eq.empty and "daily_return" in eq.columns:
        returns = eq["daily_return"].dropna()
        downside = returns[returns < 0]
        if len(returns) > 1 and len(downside) > 0:
            downside_std = float(downside.std() * (252 ** 0.5))
            ann_return = risk.get("annualized_return")
            if ann_return is not None and downside_std > 0:
                sortino = float(ann_return / downside_std)

    if engine_mode == "shared_cash":
        best_symbols: list[dict[str, Any]] = []
        worst_symbols: list[dict[str, Any]] = []
        if not rt.empty and "symbol" in rt.columns:
            sym_pnl = rt.groupby("symbol")["net_pnl"].sum().dropna().sort_values()
            sym_pnl = sym_pnl[sym_pnl != 0]
            if not sym_pnl.empty:
                worst_symbols = [{"symbol": str(k), "total_return": _json_ready(v)} for k, v in sym_pnl.head(3).items()]
                best_symbols = [{"symbol": str(k), "total_return": _json_ready(v)} for k, v in sym_pnl.tail(3).items()][::-1]

        return {
            "run_id": run_id,
            "strategy_id": strategy_id,
            "strategy_name": strategy_name,
            "strategy_family": strategy_family,
            "date_start": summary_metrics.get("date_start"),
            "date_end": summary_metrics.get("date_end"),
            "symbol_count": summary_metrics.get("symbol_count"),
            "traded_symbol_count": summary_metrics.get("traded_symbol_count"),
            "total_trade_count": total_trades,
            "open_trade_count": None,
            "winning_trade_count": win_count,
            "losing_trade_count": loss_count,
            "win_rate": win_count / total_trades if total_trades > 0 else None,
            "gross_profit": float(positive_pnl.sum()) if not positive_pnl.empty else 0.0,
            "gross_loss": float(negative_pnl.sum()) if not negative_pnl.empty else 0.0,
            "net_profit": float(rt["net_pnl"].sum()) if not rt.empty else 0.0,
            "profit_factor": abs(float(positive_pnl.sum()) / float(negative_pnl.sum())) if not negative_pnl.empty and negative_pnl.sum() != 0 else None,
            "expectancy": float(rt["net_pnl"].mean()) if not rt.empty else None,
            "recovery_factor": risk.get("return_drawdown_ratio"),
            "sharpe": summary_metrics.get("sharpe") or risk.get("sharpe"),
            "sortino": sortino,
            "information_ratio": None,
            "strategy_total_return": summary_metrics.get("total_return") or risk.get("total_return"),
            "buy_hold_return": None,
            "excess_return_vs_buy_hold": None,
            "positive_symbol_rate": None,
            "mean_symbol_return": None,
            "median_symbol_return": None,
            "max_drawdown": summary_metrics.get("max_drawdown") or risk.get("max_drawdown"),
            "max_drawdown_duration_days": summary_metrics.get("max_drawdown_duration_days") or risk.get("max_drawdown_duration_days"),
            "total_cost": summary_metrics.get("total_cost"),
            "largest_winning_trade": float(positive_pnl.max()) if not positive_pnl.empty else None,
            "largest_losing_trade": float(negative_pnl.min()) if not negative_pnl.empty else None,
            "average_winning_trade": float(positive_pnl.mean()) if not positive_pnl.empty else None,
            "average_losing_trade": float(negative_pnl.mean()) if not negative_pnl.empty else None,
            "average_win_loss_ratio": _safe_div(float(positive_pnl.mean()), abs(float(negative_pnl.mean()))) if not positive_pnl.empty and not negative_pnl.empty else None,
            "max_consecutive_wins": _max_consecutive(rt, "net_pnl", "win"),
            "max_consecutive_losses": _max_consecutive(rt, "net_pnl", "loss"),
            "average_bars_in_trade": float(rt["holding_days"].mean()) if not rt.empty else None,
            "average_bars_in_winning_trades": float(positive_pnl.index.to_series().map(
                lambda i: rt.loc[i, "holding_days"] if i in rt.index else None
            ).mean()) if not positive_pnl.empty and "holding_days" in rt.columns else None,
            "average_bars_in_losing_trades": float(negative_pnl.index.to_series().map(
                lambda i: rt.loc[i, "holding_days"] if i in rt.index else None
            ).mean()) if not negative_pnl.empty and "holding_days" in rt.columns else None,
            "best_symbols": best_symbols,
            "worst_symbols": worst_symbols,
        }

    # batch mode
    result: dict[str, Any] = {
        "run_id": run_id,
        "strategy_id": strategy_id,
        "strategy_name": strategy_name,
        "strategy_family": strategy_family,
        "date_start": summary_metrics.get("date_start"),
        "date_end": summary_metrics.get("date_end"),
        "symbol_count": summary_metrics.get("symbol_count", len(summary) if summary is not None else 0),
        "traded_symbol_count": summary_metrics.get("traded_symbol_count"),
        "total_trade_count": total_trades,
        "open_trade_count": None,
        "winning_trade_count": win_count,
        "losing_trade_count": loss_count,
        "win_rate": win_count / total_trades if total_trades > 0 else None,
        "gross_profit": float(positive_pnl.sum()) if not positive_pnl.empty else 0.0,
        "gross_loss": float(negative_pnl.sum()) if not negative_pnl.empty else 0.0,
        "net_profit": float(rt["net_pnl"].sum()) if not rt.empty else 0.0,
        "profit_factor": abs(float(positive_pnl.sum()) / float(negative_pnl.sum())) if not negative_pnl.empty and negative_pnl.sum() != 0 else None,
        "expectancy": float(rt["net_pnl"].mean()) if not rt.empty else None,
        "recovery_factor": risk.get("return_drawdown_ratio"),
        "sharpe": risk.get("sharpe"),
        "sortino": sortino,
        "information_ratio": None,
        "strategy_total_return": risk.get("total_return"),
        "buy_hold_return": summary_metrics.get("buy_hold_return"),
        "excess_return_vs_buy_hold": summary_metrics.get("excess_return_vs_buy_hold"),
        "positive_symbol_rate": summary_metrics.get("positive_return_symbol_rate"),
        "mean_symbol_return": summary_metrics.get("mean_total_return"),
        "median_symbol_return": summary_metrics.get("median_total_return"),
        "max_drawdown": risk.get("max_drawdown"),
        "max_drawdown_duration_days": risk.get("max_drawdown_duration_days"),
        "total_cost": summary_metrics.get("total_cost"),
        "largest_winning_trade": float(positive_pnl.max()) if not positive_pnl.empty else None,
        "largest_losing_trade": float(negative_pnl.min()) if not negative_pnl.empty else None,
        "average_winning_trade": float(positive_pnl.mean()) if not positive_pnl.empty else None,
        "average_losing_trade": float(negative_pnl.mean()) if not negative_pnl.empty else None,
        "average_win_loss_ratio": _safe_div(float(positive_pnl.mean()), abs(float(negative_pnl.mean()))) if not positive_pnl.empty and not negative_pnl.empty else None,
        "max_consecutive_wins": _max_consecutive(rt, "net_pnl", "win"),
        "max_consecutive_losses": _max_consecutive(rt, "net_pnl", "loss"),
        "average_bars_in_trade": float(rt["holding_days"].mean()) if not rt.empty else None,
        "average_bars_in_winning_trades": float(positive_pnl.index.to_series().map(
            lambda i: rt.loc[i, "holding_days"] if i in rt.index else None
        ).mean()) if not positive_pnl.empty and "holding_days" in rt.columns else None,
        "average_bars_in_losing_trades": float(negative_pnl.index.to_series().map(
            lambda i: rt.loc[i, "holding_days"] if i in rt.index else None
        ).mean()) if not negative_pnl.empty and "holding_days" in rt.columns else None,
        "best_symbols": _best_symbols(summary, "total_return", 3) if summary is not None else [],
        "worst_symbols": _worst_symbols(summary, "total_return", 3) if summary is not None else [],
    }
    for key, value in summary_metrics.items():
        if key.startswith("benchmark_") and key not in result:
            result[key] = value
    return result


def build_trade_analysis(round_trips: pd.DataFrame) -> dict[str, Any]:
    if round_trips.empty:
        return {
            "pnl_distribution_bins": [],
            "win_loss_counts": {"win": 0, "loss": 0},
            "percent_profitable": None,
            "average_pnl": None,
            "average_winning_trade": None,
            "average_losing_trade": None,
            "average_win_loss_ratio": None,
            "largest_winning_trade": None,
            "largest_losing_trade": None,
            "outlier_count": 0,
            "outlier_pnl": 0.0,
            "average_bars_in_trade": None,
            "average_bars_in_winning_trades": None,
            "average_bars_in_losing_trades": None,
        }

    pnl = round_trips["net_pnl"].dropna()
    positive = pnl[pnl > 0]
    negative = pnl[pnl < 0]
    total = len(pnl)
    wins = len(positive)
    losses = len(negative)

    bins: list[dict[str, Any]] = []
    if total > 0:
        pnl_min = float(pnl.min())
        pnl_max = float(pnl.max())
        bin_count = min(10, total)
        if pnl_max > pnl_min and bin_count > 0:
            bin_edges = np.linspace(pnl_min, pnl_max, bin_count + 1)
            bin_counts, _ = np.histogram(pnl, bins=bin_edges)
            for i in range(bin_count):
                bins.append({
                    "label": f"{bin_edges[i]:.2f} – {bin_edges[i + 1]:.2f}",
                    "count": int(bin_counts[i]),
                })

    mean_val = float(pnl.mean())
    std_val = float(pnl.std()) if len(pnl) > 1 else 0.0
    outlier_threshold = mean_val + 3 * std_val if std_val > 0 else None
    outliers = pnl[pnl > outlier_threshold] if outlier_threshold is not None else pd.Series(dtype=float)

    return {
        "pnl_distribution_bins": bins,
        "win_loss_counts": {"win": wins, "loss": losses},
        "percent_profitable": wins / total if total > 0 else None,
        "average_pnl": mean_val if total > 0 else None,
        "average_winning_trade": float(positive.mean()) if not positive.empty else None,
        "average_losing_trade": float(negative.mean()) if not negative.empty else None,
        "average_win_loss_ratio": _safe_div(float(positive.mean()), abs(float(negative.mean()))) if not positive.empty and not negative.empty else None,
        "largest_winning_trade": float(positive.max()) if not positive.empty else None,
        "largest_losing_trade": float(negative.min()) if not negative.empty else None,
        "outlier_count": len(outliers),
        "outlier_pnl": float(outliers.sum()),
        "average_bars_in_trade": float(round_trips["holding_days"].mean()) if "holding_days" in round_trips.columns else None,
        "average_bars_in_winning_trades": float(round_trips.loc[positive.index, "holding_days"].mean()) if not positive.empty and "holding_days" in round_trips.columns else None,
        "average_bars_in_losing_trades": float(round_trips.loc[negative.index, "holding_days"].mean()) if not negative.empty and "holding_days" in round_trips.columns else None,
    }


def build_capital_usage(
    equity_curve: pd.DataFrame,
    strategy_perf: dict[str, Any],
) -> dict[str, Any]:
    cagr = None
    if not equity_curve.empty and "total_equity" in equity_curve.columns:
        start_eq = float(equity_curve["total_equity"].iloc[0])
        end_eq = float(equity_curve["total_equity"].iloc[-1])
        if start_eq > 0:
            periods = max(len(equity_curve) - 1, 1)
            years = periods / 252.0
            if years > 0 and end_eq > 0:
                cagr = float((end_eq / start_eq) ** (1.0 / years) - 1.0)

    return {
        "cagr": cagr,
        "initial_capital_return": strategy_perf.get("strategy_total_return"),
        "required_account_size": None,
        "account_size_return": None,
        "net_profit_as_pct_of_max_loss": _safe_div(
            strategy_perf.get("net_profit"),
            abs(strategy_perf.get("max_drawdown") or 0.0),
        ),
        "average_margin_used": None,
        "maximum_margin_used": None,
        "margin_efficiency": None,
        "margin_call_count": 0,
    }


def generate_equity_curve_chart(equity: pd.DataFrame, output_path: Path) -> Path | None:
    if equity.empty or "trade_date" not in equity.columns or "total_equity" not in equity.columns:
        return None

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    import matplotlib.ticker as mticker

    dates = pd.to_datetime(equity["trade_date"])
    total_eq = equity["total_equity"]
    cum_return = equity.get("cumulative_return", pd.Series([0.0] * len(equity)))

    fig, ax1 = plt.subplots(figsize=(11.69, 4.0))
    ax1.plot(dates, total_eq, color="#1f77b4", linewidth=1.0, label="Total Equity")
    ax1.set_ylabel("Total Equity", color="#1f77b4")
    ax1.tick_params(axis="y", labelcolor="#1f77b4")
    ax1.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{x:,.0f}"))

    ax2 = ax1.twinx()
    ax2.plot(dates, cum_return, color="#ff7f0e", linewidth=1.0, label="Cumulative Return")
    ax2.set_ylabel("Cumulative Return", color="#ff7f0e")
    ax2.tick_params(axis="y", labelcolor="#ff7f0e")
    ax2.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{x:.1%}"))
    ax2.axhline(y=0, color="#ff7f0e", linewidth=0.5, linestyle="--", alpha=0.5)

    ax1.set_title("Full-Market Equity Curve")
    ax1.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax1.xaxis.set_major_locator(mdates.AutoDateLocator())
    fig.autofmt_xdate(rotation=30)
    ax1.grid(True, alpha=0.3)

    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc="upper left")

    fig.tight_layout()
    fig.savefig(str(output_path), dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output_path


def generate_drawdown_chart(drawdown: pd.DataFrame, output_path: Path) -> Path | None:
    if drawdown.empty or "trade_date" not in drawdown.columns or "drawdown_pct" not in drawdown.columns:
        return None

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    import matplotlib.ticker as mticker

    dates = pd.to_datetime(drawdown["trade_date"])
    dd_pct = drawdown["drawdown_pct"]

    fig, ax = plt.subplots(figsize=(11.69, 3.2))
    ax.fill_between(dates, 0, dd_pct, color="#d62728", alpha=0.25, label="Drawdown")
    ax.plot(dates, dd_pct, color="#d62728", linewidth=0.8)
    ax.axhline(y=0, color="#333333", linewidth=0.5, linestyle="--")

    ax.set_ylabel("Drawdown")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{x:.1%}"))
    ax.set_title("Drawdown Curve")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax.xaxis.set_major_locator(mdates.AutoDateLocator())
    fig.autofmt_xdate(rotation=30)
    ax.grid(True, alpha=0.3)
    ax.legend(loc="lower left")

    fig.tight_layout()
    fig.savefig(str(output_path), dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output_path


def write_full_market_excel(
    output_dir: Path,
    overview: dict[str, Any],
    trades_ledger: pd.DataFrame,
    round_trips: pd.DataFrame,
    symbol_perf: pd.DataFrame,
    equity: pd.DataFrame,
    drawdown: pd.DataFrame,
    trade_analysis: dict[str, Any],
    capital_usage: dict[str, Any],
    *,
    strategy_perf: dict[str, Any] | None = None,
    equity_curve_chart: Path | None = None,
    drawdown_curve_chart: Path | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    header_font = Font(bold=True)
    section_font = Font(bold=True)

    def _write_df_sheet(ws, df: pd.DataFrame) -> None:
        ws.append([COLUMN_LABELS.get(str(c), str(c)) for c in df.columns])
        for cell in ws[1]:
            cell.font = header_font
        number_formats = [_column_number_format(str(c), df[c]) for c in df.columns]
        for _, row in df.iterrows():
            ws.append([_excel_value(v) for v in row])
        for col_idx, fmt in enumerate(number_formats, start=1):
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = fmt

    def _write_kv_sheet(ws, data: dict[str, Any]) -> None:
        ws.append(["项目", "值"])
        for cell in ws[1]:
            cell.font = header_font
        row_formats: dict[int, str] = {}
        for key, value in data.items():
            label = _key_label(key)
            if isinstance(value, dict):
                # Nested block (e.g. 策略说明): section row then its items.
                ws.append([label, ""])
                ws.cell(row=ws.max_row, column=1).font = section_font
                for sub_key, sub_value in value.items():
                    sub_label = _key_label(sub_key)
                    ws.append([sub_label, _excel_value(sub_value)])
                    row_formats[ws.max_row] = _column_number_format(str(sub_key))
                continue
            ws.append([label, _excel_value(value)])
            row_formats[ws.max_row] = _column_number_format(str(key))
        for row_idx, fmt in row_formats.items():
            cell = ws.cell(row=row_idx, column=2)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = fmt

    # Overview
    _write_kv_sheet(wb.active, overview)
    wb.active.title = "Overview"

    # Performance
    ws_perf = wb.create_sheet("Performance")
    _write_kv_sheet(ws_perf, strategy_perf or {})

    # TradesLedger
    ws_ledger = wb.create_sheet("TradesLedger")
    _write_df_sheet(ws_ledger, trades_ledger)

    # RoundTrips
    ws_rt = wb.create_sheet("RoundTrips")
    _write_df_sheet(ws_rt, round_trips)

    # SymbolPerformance
    ws_sp = wb.create_sheet("SymbolPerformance")
    _write_df_sheet(ws_sp, symbol_perf)

    # EquityCurve
    ws_eq = wb.create_sheet("EquityCurve")
    _write_df_sheet(ws_eq, equity)

    # DrawdownCurve
    ws_dd = wb.create_sheet("DrawdownCurve")
    _write_df_sheet(ws_dd, drawdown)

    # TradeAnalysis
    ws_ta = wb.create_sheet("TradeAnalysis")
    _write_kv_sheet(ws_ta, trade_analysis)

    # CapitalUsage
    ws_cu = wb.create_sheet("CapitalUsage")
    _write_kv_sheet(ws_cu, capital_usage)

    # EquityCurveChart (embedded image)
    if equity_curve_chart is not None and equity_curve_chart.exists():
        from openpyxl.drawing.image import Image as XlImage
        ws_chart = wb.create_sheet("EquityCurveChart")
        img = XlImage(str(equity_curve_chart))
        img.width = 820
        img.height = 310
        ws_chart.add_image(img, "A1")

    # DrawdownCurveChart (embedded image)
    if drawdown_curve_chart is not None and drawdown_curve_chart.exists():
        from openpyxl.drawing.image import Image as XlImage
        ws_chart = wb.create_sheet("DrawdownCurveChart")
        img = XlImage(str(drawdown_curve_chart))
        img.width = 820
        img.height = 250
        ws_chart.add_image(img, "A1")

    # Artifacts
    ws_art = wb.create_sheet("Artifacts")
    ws_art.append(["文件", "说明"])
    for cell in ws_art[1]:
        cell.font = header_font
    for name, desc in _ARTIFACT_DESCRIPTIONS:
        ws_art.append([name, desc])

    path = output_dir / "backtest_report.xlsx"
    wb.save(str(path))
    return path


def generate_html_report(
    output_dir: Path,
    overview: dict[str, Any],
    strategy_perf: dict[str, Any],
    trades_ledger: pd.DataFrame,
    round_trips: pd.DataFrame,
    symbol_perf: pd.DataFrame,
    equity: pd.DataFrame,
    drawdown: pd.DataFrame,
    trade_analysis: dict[str, Any],
    capital_usage: dict[str, Any],
    *,
    equity_curve_chart: Path | None = None,
    drawdown_curve_chart: Path | None = None,
) -> str:
    """Generate a self-contained HTML backtest report string.

    Charts are embedded as base64 data URIs. The output is a single,
    offline-browsable HTML document with inline CSS.
    """
    ec_b64 = _chart_base64(equity_curve_chart) if equity_curve_chart else ""
    dd_b64 = _chart_base64(drawdown_curve_chart) if drawdown_curve_chart else ""

    def _kv_table(data: dict[str, Any], columns: tuple[str, str] = ("项目", "值")) -> str:
        lines = ['<table class="kv">', f"<thead><tr><th>{columns[0]}</th><th>{columns[1]}</th></tr></thead><tbody>"]
        for k, v in data.items():
            label = _key_label(k)
            if isinstance(v, dict):
                lines.append(f'<tr><td colspan="2"><strong>{_h(label)}</strong></td></tr>')
                for sk, sv in v.items():
                    sub_label = _key_label(sk)
                    lines.append(f"<tr><td>{_h(sub_label)}</td><td>{_h(_fmt_display(sv, ratio=is_ratio_column(str(sk))))}</td></tr>")
                continue
            lines.append(f"<tr><td>{_h(label)}</td><td>{_h(_fmt_display(v, ratio=is_ratio_column(str(k))))}</td></tr>")
        lines.append("</tbody></table>")
        return "\n".join(lines)

    def _df_table(df: pd.DataFrame, max_rows: int = 200) -> str:
        if df.empty:
            return '<p class="empty">暂无数据</p>'
        cols = [str(c) for c in df.columns]
        ratio_flags = {c: is_ratio_column(c, df[c]) for c in df.columns}
        lines = ['<div class="table-wrap"><table>', "<thead><tr>"]
        lines.extend(f"<th>{_h(COLUMN_LABELS.get(c, c))}</th>" for c in cols)
        lines.append("</tr></thead><tbody>")
        for i, (_, row) in enumerate(df.iterrows()):
            if i >= max_rows:
                lines.append(f'<tr><td colspan="{len(cols)}" class="truncated">... 其余 {len(df) - max_rows} 行未显示</td></tr>')
                break
            lines.append("<tr>")
            for c in cols:
                lines.append(f"<td>{_h(_fmt_display(row[c], ratio=ratio_flags[c]))}</td>")
            lines.append("</tr>")
        lines.append("</tbody></table></div>")
        return "\n".join(lines)

    def _metrics_cards(perf: dict[str, Any]) -> str:
        cards = [
            ("总收益率", _fmt_pct(perf.get("strategy_total_return"))),
            ("胜率", _fmt_pct(perf.get("win_rate"))),
            ("夏普比率", _fmt_num(perf.get("sharpe"), 2)),
            ("最大回撤", _fmt_pct(perf.get("max_drawdown"))),
            ("交易笔数", str(perf.get("total_trade_count") or "-")),
            ("盈亏比", _fmt_num(perf.get("profit_factor"), 2)),
        ]
        lines = ['<div class="metrics-grid">']
        for label, value in cards:
            lines.append(f'<div class="metric-card"><div class="metric-label">{label}</div><div class="metric-value">{value}</div></div>')
        lines.append("</div>")
        return "\n".join(lines)

    def _pnl_distribution(ta: dict[str, Any]) -> str:
        bins = ta.get("pnl_distribution_bins") or []
        if not bins:
            return ""
        max_count = max(b["count"] for b in bins) if bins else 1
        lines = ['<div class="pnl-bars">']
        for b in bins:
            w = int(b["count"] / max_count * 100) if max_count > 0 else 0
            lines.append(
                f'<div class="pnl-bar-row">'
                f'<span class="pnl-label">{_h(b["label"])}</span>'
                f'<span class="pnl-bar"><span class="pnl-fill" style="width:{w}%"></span></span>'
                f'<span class="pnl-count">{b["count"]}</span>'
                f'</div>'
            )
        lines.append("</div>")
        return "\n".join(lines)

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>回测报告 — {_h(str(overview.get('strategy_name') or overview.get('run_id') or ''))}</title>
<style>
  :root {{--bg:#fff;--fg:#1a1a1a;--muted:#666;--border:#e0e0e0;--accent:#1f77b4;--green:#2ca02c;--red:#d62728;--card-bg:#f8f9fa;--radius:6px}}
  * {{box-sizing:border-box;margin:0;padding:0}}
  body {{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;color:var(--fg);background:var(--bg);line-height:1.5;padding:0}}
  .container {{max-width:1200px;margin:0 auto;padding:24px}}
  header {{background:linear-gradient(135deg,#1a1a2e,#16213e);color:#fff;padding:32px 24px;text-align:center}}
  header h1 {{font-size:1.6rem;font-weight:600;margin:0 0 8px}}
  header .meta {{font-size:0.85rem;opacity:0.8}}
  h2 {{font-size:1.2rem;font-weight:600;margin:32px 0 16px;padding-bottom:8px;border-bottom:2px solid var(--border);color:#1a1a2e}}
  h3 {{font-size:1rem;font-weight:600;margin:20px 0 12px}}
  table {{width:100%;border-collapse:collapse;font-size:0.875rem}}
  th,td {{padding:8px 12px;text-align:left;border-bottom:1px solid var(--border)}}
  th {{background:var(--card-bg);font-weight:600;white-space:nowrap;position:sticky;top:0}}
  tr:hover td {{background:#f0f4ff}}
  .table-wrap {{overflow-x:auto;max-height:500px;overflow-y:auto;border:1px solid var(--border);border-radius:var(--radius)}}
  .kv {{max-width:800px}}
  .kv td:first-child {{font-weight:500;color:var(--muted);width:240px}}
  .metrics-grid {{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin-bottom:24px}}
  .metric-card {{background:var(--card-bg);border:1px solid var(--border);border-radius:var(--radius);padding:16px;text-align:center}}
  .metric-label {{font-size:0.78rem;color:var(--muted);margin-bottom:4px}}
  .metric-value {{font-size:1.3rem;font-weight:700}}
  .chart {{text-align:center;margin:16px 0}}
  .chart img {{max-width:100%;border:1px solid var(--border);border-radius:var(--radius)}}
  .pnl-bars {{margin:12px 0}}
  .pnl-bar-row {{display:flex;align-items:center;gap:8px;margin:4px 0;font-size:0.8rem}}
  .pnl-label {{width:240px;text-align:right;color:var(--muted);white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}
  .pnl-bar {{flex:1;height:18px;background:#e8e8e8;border-radius:3px;overflow:hidden}}
  .pnl-fill {{display:block;height:100%;background:var(--accent);border-radius:3px;min-width:1px}}
  .pnl-count {{width:36px;text-align:left;font-weight:600}}
  .section {{margin-bottom:16px}}
  .positive {{color:var(--green)}}
  .negative {{color:var(--red)}}
  .empty {{color:var(--muted);font-style:italic;padding:12px}}
  .truncated {{color:var(--muted);font-style:italic;text-align:center}}
  footer {{text-align:center;color:var(--muted);font-size:0.8rem;padding:24px;border-top:1px solid var(--border);margin-top:32px}}
  @media (max-width:768px) {{.container{{padding:12px}} .metrics-grid{{grid-template-columns:repeat(2,1fr)}} .kv td:first-child{{width:140px}}}}
</style>
</head>
<body>
<header>
  <h1>回测报告</h1>
  <div class="meta">
    {_h(str(overview.get('strategy_name') or ''))} &nbsp;|&nbsp;
    引擎: {_h(str(overview.get('engine_mode', '')))} &nbsp;|&nbsp;
    {_h(str(overview.get('market', '')))} &nbsp;|&nbsp;
    运行: {_h(str(overview.get('run_id') or ''))}
  </div>
</header>
<div class="container">

<h2>0. 策略说明</h2>
{_kv_table({'strategy_intro': overview.get('strategy_intro') or {}})}

<h2>1. 绩效汇总</h2>
{_metrics_cards(strategy_perf)}
{_kv_table(strategy_perf)}

<h2>2. 图表</h2>
<div class="chart"><h3>净值曲线</h3>{f'<img src="data:image/png;base64,{ec_b64}" alt="净值曲线">' if ec_b64 else '<p class="empty">暂无净值数据</p>'}</div>
<div class="chart"><h3>回撤曲线</h3>{f'<img src="data:image/png;base64,{dd_b64}" alt="回撤曲线">' if dd_b64 else '<p class="empty">暂无回撤数据</p>'}</div>

<h2>3. 交易分析</h2>
<h3>盈亏分布</h3>
{_pnl_distribution(trade_analysis)}
<div class="section">{_kv_table(trade_analysis)}</div>

<h2>4. 个股表现</h2>
{_df_table(symbol_perf, max_rows=100)}

<h2>5. 资金使用</h2>
{_kv_table(capital_usage)}

<h2>6. 成交流水</h2>
{_df_table(trades_ledger, max_rows=200)}

<h2>7. 完整回合</h2>
{_df_table(round_trips, max_rows=200)}

<h2>8. 运行概览</h2>
{_kv_table(overview)}

</div>
<footer>生成时间 {overview.get('generated_at', 'N/A')} &mdash; quant-research-engine</footer>
</body>
</html>"""
    return html


def _chart_base64(path: Path) -> str:
    if not path.exists():
        return ""
    return base64.b64encode(path.read_bytes()).decode("ascii")


def _h(text: str) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _fmt_pct(value: Any) -> str:
    if _is_na(value):
        return "-"
    try:
        f = float(value)
        return f"{f:.2%}"
    except (ValueError, TypeError):
        return str(value)


def _fmt_num(value: Any, decimals: int = 2) -> str:
    if _is_na(value):
        return "-"
    try:
        f = float(value)
        return f"{f:.{decimals}f}"
    except (ValueError, TypeError):
        return str(value)


def write_full_market_report(
    result: Any,
    output_dir: str | Path,
    run_request: dict[str, Any] | None = None,
    run_id: str | None = None,
    engine_mode: str = "batch",
) -> dict[str, Path]:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)

    strategy_id = None
    strategy_name = None
    strategy_family = None
    req: dict[str, Any] = {}
    if run_request is not None:
        req = run_request.get("request", {}) or {}
        strategy_id = req.get("strategy_id")
        strategy_name = req.get("strategy_name")
        strategy_family = req.get("strategy_family")
    strategy_intro = _build_strategy_intro(run_request)

    if engine_mode == "shared_cash":
        orders = getattr(result, "orders", pd.DataFrame())
        trades = getattr(result, "trades", pd.DataFrame())
        eq_curve = getattr(result, "equity_curve", pd.DataFrame())
        shared_summary = getattr(result, "summary", {})

        trades_ledger = build_trades_ledger(orders, run_id=run_id, strategy_id=strategy_id, engine_mode="shared_cash")
        round_trips = build_round_trips(trades, run_id=run_id, strategy_id=strategy_id, engine_mode="shared_cash")
        symbol_perf = build_symbol_performance(
            engine_mode="shared_cash", round_trips=round_trips, orders=orders,
            failed_orders=getattr(result, "failed_orders", pd.DataFrame()),
        )
        fm_equity_curve = build_full_market_equity_curve(eq_curve, engine_mode="shared_cash")
        drawdown_curve = build_drawdown_curve(fm_equity_curve)
        strategy_perf = build_strategy_performance(
            shared_summary, round_trips=round_trips, equity_curve=fm_equity_curve,
            run_id=run_id, strategy_id=strategy_id,
            strategy_name=strategy_name, strategy_family=strategy_family,
            engine_mode="shared_cash",
        )
        if run_request is not None:
            _req = run_request.get("request", {})
            if strategy_perf.get("date_start") is None:
                strategy_perf["date_start"] = _req.get("date_start") or run_request.get("start")
            if strategy_perf.get("date_end") is None:
                strategy_perf["date_end"] = _req.get("date_end") or run_request.get("end")
        trade_analysis = build_trade_analysis(round_trips)
        capital_usage = build_capital_usage(fm_equity_curve, strategy_perf)

        equity_curve_chart_path = generate_equity_curve_chart(fm_equity_curve, directory / "equity_curve_chart.png")
        drawdown_curve_chart_path = generate_drawdown_chart(drawdown_curve, directory / "drawdown_curve_chart.png")

        symbol_count = 0
        if not orders.empty and "symbol" in orders.columns:
            symbol_count = orders["symbol"].nunique()

        overview: dict[str, Any] = {
            "run_id": run_id,
            "strategy_id": strategy_id,
            "strategy_name": strategy_name,
            "strategy_family": strategy_family,
            "strategy_description": None,
            "strategy_intro": strategy_intro,
            "engine_mode": "shared_cash",
            "date_start": shared_summary.get("date_start") or req.get("date_start"),
            "date_end": shared_summary.get("date_end") or req.get("date_end"),
            "market": "ashare",
            "frequency": "daily",
            "universe": "all_a",
            "symbol_count": shared_summary.get("symbol_count", symbol_count),
            "traded_symbol_count": shared_summary.get("traded_symbol_count"),
            "initial_cash": shared_summary.get("initial_cash"),
            "market_profile": shared_summary.get("market_profile"),
            "entry_execution": "t_plus_1_open",
            "exit_execution": "t_plus_1_open",
            "cost_model": "ashare_default",
            "slippage_model": "percent",
            "factor_requirements": None,
            "strategy_source_path": run_request.get("strategy_path") if run_request else None,
            "generated_at": pd.Timestamp.now(tz="UTC").isoformat(),
        }

        paths: dict[str, Path] = {
            "trades_ledger": directory / "trades_ledger.csv",
            "round_trips": directory / "round_trips.csv",
            "symbol_performance": directory / "symbol_performance.csv",
            "strategy_performance": directory / "strategy_performance.json",
            "full_market_equity_curve": directory / "full_market_equity_curve.csv",
            "drawdown_curve": directory / "drawdown_curve.csv",
            "trade_analysis": directory / "trade_analysis.json",
            "capital_usage": directory / "capital_usage.json",
        }
        if equity_curve_chart_path is not None:
            paths["equity_curve_chart"] = equity_curve_chart_path
        if drawdown_curve_chart_path is not None:
            paths["drawdown_curve_chart"] = drawdown_curve_chart_path
        paths["backtest_report"] = write_full_market_excel(
            directory, overview, trades_ledger, round_trips, symbol_perf,
            fm_equity_curve, drawdown_curve, trade_analysis, capital_usage,
            strategy_perf=strategy_perf,
            equity_curve_chart=equity_curve_chart_path,
            drawdown_curve_chart=drawdown_curve_chart_path,
        )

        trades_ledger.to_csv(paths["trades_ledger"], index=False)
        round_trips.to_csv(paths["round_trips"], index=False)
        symbol_perf.to_csv(paths["symbol_performance"], index=False)
        fm_equity_curve.to_csv(paths["full_market_equity_curve"], index=False)
        drawdown_curve.to_csv(paths["drawdown_curve"], index=False)
        paths["strategy_performance"].write_text(
            json.dumps(_json_ready(strategy_perf), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        paths["trade_analysis"].write_text(
            json.dumps(_json_ready(trade_analysis), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        paths["capital_usage"].write_text(
            json.dumps(_json_ready(capital_usage), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        paths["html_report"] = directory / "backtest_report.html"
        paths["html_report"].write_text(
            generate_html_report(
                directory, overview, strategy_perf,
                trades_ledger, round_trips, symbol_perf,
                fm_equity_curve, drawdown_curve,
                trade_analysis, capital_usage,
                equity_curve_chart=equity_curve_chart_path,
                drawdown_curve_chart=drawdown_curve_chart_path,
            ),
            encoding="utf-8",
        )

        return paths

    # batch mode
    summary = getattr(result, "summary", pd.DataFrame())
    summary_metrics = getattr(result, "summary_metrics", {})
    results = getattr(result, "results", {})

    trades_ledger = build_trades_ledger(results, run_id=run_id, strategy_id=strategy_id)
    round_trips = build_round_trips(results, run_id=run_id, strategy_id=strategy_id)
    symbol_perf = build_symbol_performance(summary, results) if not summary.empty else pd.DataFrame(columns=_SYMBOL_PERFORMANCE_COLUMNS)
    equity_curve = build_full_market_equity_curve(results)
    drawdown_curve = build_drawdown_curve(equity_curve)
    strategy_perf = build_strategy_performance(
        summary_metrics, summary, round_trips, equity_curve,
        run_id=run_id, strategy_id=strategy_id,
        strategy_name=strategy_name, strategy_family=strategy_family,
    )
    if run_request is not None:
        # Backfill run-window fields the engine summary does not carry
        _req = run_request.get("request", {})
        if strategy_perf.get("date_start") is None:
            strategy_perf["date_start"] = _req.get("date_start") or run_request.get("start")
        if strategy_perf.get("date_end") is None:
            strategy_perf["date_end"] = _req.get("date_end") or run_request.get("end")
    trade_analysis = build_trade_analysis(round_trips)
    capital_usage = build_capital_usage(equity_curve, strategy_perf)

    equity_curve_chart_path = generate_equity_curve_chart(equity_curve, directory / "equity_curve_chart.png")
    drawdown_curve_chart_path = generate_drawdown_chart(drawdown_curve, directory / "drawdown_curve_chart.png")

    overview: dict[str, Any] = {
        "run_id": run_id,
        "strategy_id": strategy_id,
        "strategy_name": strategy_name,
        "strategy_family": strategy_family,
        "strategy_description": None,
        "strategy_intro": strategy_intro,
        "engine_mode": "batch",
        "date_start": summary_metrics.get("date_start") or req.get("date_start"),
        "date_end": summary_metrics.get("date_end") or req.get("date_end"),
        "market": "ashare",
        "frequency": "daily",
        "universe": "all_a",
        "symbol_count": summary_metrics.get("symbol_count"),
        "traded_symbol_count": summary_metrics.get("traded_symbol_count"),
        "initial_cash": req.get("initial_cash"),
        "market_profile": summary_metrics.get("market_profile"),
        "entry_execution": "t_plus_1_open",
        "exit_execution": "t_plus_1_open",
        "cost_model": "ashare_default",
        "slippage_model": "percent",
        "factor_requirements": None,
        "strategy_source_path": run_request.get("strategy_path") if run_request else None,
        "generated_at": pd.Timestamp.now(tz="UTC").isoformat(),
    }

    paths: dict[str, Path] = {
        "trades_ledger": directory / "trades_ledger.csv",
        "round_trips": directory / "round_trips.csv",
        "symbol_performance": directory / "symbol_performance.csv",
        "strategy_performance": directory / "strategy_performance.json",
        "equity_curve": directory / "equity_curve.csv",
        "drawdown_curve": directory / "drawdown_curve.csv",
        "trade_analysis": directory / "trade_analysis.json",
        "capital_usage": directory / "capital_usage.json",
    }
    if equity_curve_chart_path is not None:
        paths["equity_curve_chart"] = equity_curve_chart_path
    if drawdown_curve_chart_path is not None:
        paths["drawdown_curve_chart"] = drawdown_curve_chart_path
    paths["backtest_report"] = write_full_market_excel(
        directory, overview, trades_ledger, round_trips, symbol_perf,
        equity_curve, drawdown_curve, trade_analysis, capital_usage,
        strategy_perf=strategy_perf,
        equity_curve_chart=equity_curve_chart_path,
        drawdown_curve_chart=drawdown_curve_chart_path,
    )

    # Write CSV/JSON artifacts
    trades_ledger.to_csv(paths["trades_ledger"], index=False)
    round_trips.to_csv(paths["round_trips"], index=False)
    symbol_perf.to_csv(paths["symbol_performance"], index=False)
    equity_curve.to_csv(paths["equity_curve"], index=False)
    drawdown_curve.to_csv(paths["drawdown_curve"], index=False)
    paths["strategy_performance"].write_text(
        json.dumps(_json_ready(strategy_perf), ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    paths["trade_analysis"].write_text(
        json.dumps(_json_ready(trade_analysis), ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    paths["capital_usage"].write_text(
        json.dumps(_json_ready(capital_usage), ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    paths["html_report"] = directory / "backtest_report.html"
    paths["html_report"].write_text(
        generate_html_report(
            directory, overview, strategy_perf,
            trades_ledger, round_trips, symbol_perf,
            equity_curve, drawdown_curve,
            trade_analysis, capital_usage,
            equity_curve_chart=equity_curve_chart_path,
            drawdown_curve_chart=drawdown_curve_chart_path,
        ),
        encoding="utf-8",
    )

    return paths


# ---- helpers ----


def _to_date(value: Any) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    try:
        return pd.Timestamp(value).date().isoformat()
    except Exception:
        return str(value)


def _safe_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def _safe_div(a: float | None, b: float) -> float | None:
    if a is None or b == 0.0:
        return None
    return a / b


def _excel_value(value: Any) -> Any:
    if _is_na(value):
        return "-"
    if isinstance(value, (int, float, str, bool)):
        return value
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, list):
        return json.dumps(_json_ready(value), ensure_ascii=False)
    if isinstance(value, dict):
        return json.dumps(_json_ready(value), ensure_ascii=False)
    return str(value)


def _max_consecutive(frame: pd.DataFrame, column: str, direction: str) -> int:
    if frame.empty or column not in frame.columns:
        return 0
    max_count = 0
    current = 0
    for val in frame[column]:
        if direction == "win" and val > 0:
            current += 1
        elif direction == "loss" and val < 0:
            current += 1
        else:
            max_count = max(max_count, current)
            current = 0
    return max(max_count, current)


def _best_symbols(summary: pd.DataFrame, column: str, n: int) -> list[dict[str, Any]]:
    if summary.empty or column not in summary.columns:
        return []
    top = summary.dropna(subset=[column]).sort_values(column, ascending=False).head(n)
    return [{str(k): _json_ready(v) for k, v in record.items()} for record in top.to_dict("records")]


def _worst_symbols(summary: pd.DataFrame, column: str, n: int) -> list[dict[str, Any]]:
    if summary.empty or column not in summary.columns:
        return []
    bottom = summary.dropna(subset=[column]).sort_values(column, ascending=True).head(n)
    return [{str(k): _json_ready(v) for k, v in record.items()} for record in bottom.to_dict("records")]


def _json_ready(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {str(k): _json_ready(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_json_ready(v) for v in obj]
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        if pd.isna(obj):
            return None
        return float(obj)
    if isinstance(obj, (pd.Timestamp,)):
        return obj.isoformat()
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    return obj
